#!/usr/bin/env python3
"""Generate a theme's base.pptx (slide master and layouts) from theme.yaml, under slides:.

    imprint make-base-pptx --theme default
    imprint make-base-pptx --theme /path/to/brand/themes/my-theme

The master is built from code, not exported from PowerPoint, so the same theme
always gives the same master and a brand is a handful of values plus a logo:

  * 16:9, the fonts for Latin and Thai, the colour scheme (accents are what
    PowerPoint offers first for charts and shapes)
  * title and body text styles, with the bullet written on every level
  * six layouts: Title Slide, Title and Content, Title Only, Section Header,
    Two Content, Blank. The rest of PowerPoint's defaults are removed.
  * an optional logo on the master and optional line art on the title slide

Re-run it after changing anything under slides: in theme.yaml, and commit the
base.pptx it writes next to theme.yaml.
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches

from ..theme import Theme

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
KEEP = ("Title Slide", "Title and Content", "Title Only", "Section Header", "Two Content", "Blank")
ASPECT = {"16:9": (12192000, 6858000), "4:3": (9144000, 6858000)}


def _a(tag, **attrs):
    el = etree.Element(f"{{{A}}}{tag}")
    for k, v in attrs.items():
        el.set(k, str(v))
    return el


def _fill(color):
    f = _a("solidFill")
    f.append(_a("srgbClr", val=color))
    return f


def _rpr(tag, size_pt, color, font, bold=False):
    r = _a(tag, sz=int(size_pt * 100), b=int(bold), kern=1200)
    r.append(_fill(color))
    for t in ("latin", "ea", "cs"):
        r.append(_a(t, typeface=font))
    return r


def _ppr(tag, *, algn="l", mar_l=0, indent=0, bullet=None, space_before=0, line_pct=100):
    p = _a(tag, marL=mar_l, indent=indent, algn=algn, defTabSz=914400, rtl=0, eaLnBrk=1,
           latinLnBrk=0, hangingPunct=1)
    ln = _a("lnSpc"); ln.append(_a("spcPct", val=line_pct * 1000)); p.append(ln)
    sb = _a("spcBef"); sb.append(_a("spcPts", val=int(space_before * 100))); p.append(sb)
    if bullet:
        p.append(_a("buFont", typeface="Arial"))
        p.append(_a("buChar", char=bullet))
    else:
        p.append(_a("buNone"))
    return p


def _theme_part(master):
    return master.part.part_related_by(RT.THEME)


def set_theme_fonts_and_colors(master, s: dict, heading: str, body: str):
    part = _theme_part(master)
    root = etree.fromstring(part.blob)
    ns = {"a": A}
    for kind, face in (("majorFont", heading), ("minorFont", body)):
        el = root.find(f".//a:fontScheme/a:{kind}", ns)
        el.find("a:latin", ns).set("typeface", face)
        el.find("a:ea", ns).set("typeface", face)
        el.find("a:cs", ns).set("typeface", face)
        for f in el.findall("a:font", ns):
            if f.get("script") == "Thai":
                el.remove(f)
        el.append(_a("font", script="Thai", typeface=face))
    c = s.get("colors", {})
    accents = list(s.get("accents", ["2D3748", "4A5568", "718096", "A0AEC0", "2B6CB0", "1A202C"]))
    scheme = {"dk1": c.get("ink", "1A202C"), "lt1": c.get("paper", "FFFFFF"),
              "dk2": c.get("muted", "4A5568"), "lt2": c.get("rule", "E2E8F0"),
              "hlink": accents[0], "folHlink": c.get("muted", "4A5568")}
    for i, v in enumerate(accents[:6], 1):
        scheme[f"accent{i}"] = v
    clr = root.find(".//a:clrScheme", ns)
    clr.set("name", s.get("name", "Imprint"))
    for slot, value in scheme.items():
        el = clr.find(f"a:{slot}", ns)
        for child in list(el):
            el.remove(child)
        el.append(_a("srgbClr", val=str(value).lstrip("#").upper()))
    part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def set_text_styles(master, s: dict, heading: str, body: str, bullet: str):
    size = s.get("size", {})
    c = s.get("colors", {})
    ink = c.get("ink", "1A202C")
    levels = size.get("levels", [18, 16, 14])
    tx = master._element.find(qn("p:txStyles"))
    title = tx.find(qn("p:titleStyle"))
    for child in list(title):
        title.remove(child)
    p = _ppr("lvl1pPr", line_pct=100)
    p.append(_rpr("defRPr", size.get("title", 26), ink, heading, bold=True))
    title.append(p)
    bodys = tx.find(qn("p:bodyStyle"))
    for child in list(bodys):
        bodys.remove(child)
    step = 292608  # 0.32 in, the same indent the deck builder writes
    for lvl in range(1, 10):
        pt = levels[min(lvl - 1, len(levels) - 1)]
        p = _ppr(f"lvl{lvl}pPr", mar_l=step * lvl, indent=-step, bullet=bullet, space_before=6, line_pct=110)
        p.append(_rpr("defRPr", pt, ink, body))
        bodys.append(p)


def _drop_placeholders(shapes, types):
    for ph in list(shapes.placeholders):
        if ph.placeholder_format.type in types:
            ph._element.getparent().remove(ph._element)


def _place(ph, x, y, w, h):
    ph.left, ph.top, ph.width, ph.height = int(x), int(y), int(w), int(h)


def _lst_style(ph, size_pt, color, font, *, bold=False, anchor=None, algn="l"):
    txBody = ph._element.find(qn("p:txBody"))
    bodyPr = txBody.find(qn("a:bodyPr"))
    if anchor:
        bodyPr.set("anchor", anchor)
    lst = txBody.find(qn("a:lstStyle"))
    for child in list(lst):
        lst.remove(child)
    p = _ppr("lvl1pPr", algn=algn)
    p.append(_rpr("defRPr", size_pt, color, font, bold=bold))
    lst.append(p)


def _background(layout, color):
    cSld = layout._element.find(qn("p:cSld"))
    for old in cSld.findall(qn("p:bg")):
        cSld.remove(old)
    bg = etree.SubElement(cSld, f"{{{P}}}bg")
    cSld.remove(bg)
    cSld.insert(0, bg)
    bgPr = etree.SubElement(bg, f"{{{P}}}bgPr")
    bgPr.append(_fill(color))
    bgPr.append(_a("effectLst"))


def _place_picture(prs, target, path: Path, x, y, *, height=None, width=None):
    """Master and layout shape trees cannot add pictures directly: add one to a
    scratch slide, re-point its image at the target part, and move it over."""
    scratch = prs.slides.add_slide(prs.slide_layouts.get_by_name("Blank"))
    pic = scratch.shapes.add_picture(str(path), int(x), int(y), width=width, height=height)
    blip = pic._element.find(".//" + qn("a:blip"))
    image_part = scratch.part.related_part(blip.get(qn("r:embed")))
    blip.set(qn("r:embed"), target.part.relate_to(image_part, RT.IMAGE))
    target.shapes._spTree.append(pic._element)
    sldIdLst = prs.slides._sldIdLst
    last = sldIdLst[-1]
    prs.part.drop_rel(last.rId)
    sldIdLst.remove(last)
    return pic


def _place_rule(prs, target, x, y, w, h, color: str):
    """A solid bar on a master or layout, made the same way as _place_picture."""
    from pptx.enum.shapes import MSO_SHAPE
    scratch = prs.slides.add_slide(prs.slide_layouts.get_by_name("Blank"))
    bar = scratch.shapes.add_shape(MSO_SHAPE.RECTANGLE, int(x), int(y), int(w), int(h))
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor.from_string(str(color).lstrip("#").upper())
    bar.line.fill.background()
    bar.shadow.inherit = False
    target.shapes._spTree.append(bar._element)
    sldIdLst = prs.slides._sldIdLst
    last = sldIdLst[-1]
    prs.part.drop_rel(last.rId)
    sldIdLst.remove(last)
    return bar


def build(theme: Theme) -> Path:
    s = theme.slides or {}
    out = theme.source.parent / s.get("base", "base.pptx")
    font = s.get("font", {})
    heading = font.get("heading", theme.heading_font)
    body = font.get("body", theme.body_font)
    size = s.get("size", {})
    c = s.get("colors", {})
    ink, paper = c.get("ink", "1A202C"), c.get("paper", "FFFFFF")
    soft, muted = c.get("soft", "6B7280"), c.get("muted", "4A5568")
    bullet = str(s.get("bullet_char", "•"))

    prs = Presentation()
    W, H = ASPECT[str(s.get("aspect", "16:9"))]
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    M = Inches(float(s.get("margin_in", 0.6)))
    master = prs.slide_master

    for layout in list(prs.slide_layouts):
        if layout.name not in KEEP:
            prs.slide_layouts.remove(layout)

    set_theme_fonts_and_colors(master, s, heading, body)
    set_text_styles(master, s, heading, body, bullet)

    furniture = (PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER)
    _drop_placeholders(master, furniture)
    for layout in prs.slide_layouts:
        _drop_placeholders(layout, furniture)

    logo = s.get("logo")
    logo_path = (theme.source.parent / logo).resolve() if logo else None
    logo_h = Inches(float(s.get("logo_h_in", 0.3)))
    logo_w = 0
    if logo_path and logo_path.exists():
        from PIL import Image
        with Image.open(logo_path) as im:
            logo_w = int(logo_h * im.size[0] / im.size[1])
    logo_at = str(s.get("logo_position", "top-right"))
    title_top, title_h = Inches(0.45), Inches(0.95)
    rule = bool(s.get("title_rule", False))
    rule_w = Inches(float(s.get("title_rule_w_in", 0.045)))
    title_x = M
    if logo_w and logo_at == "title":
        # the logo sits left of every content title, with an optional rule between
        title_x = M + logo_w + Inches(0.25) + (rule_w + Inches(0.3) if rule else 0)
    title_w = W - title_x - M - (logo_w + Inches(0.3) if logo_w and logo_at == "top-right" else 0)
    body_top = title_top + title_h + Inches(0.2)
    body_h = H - body_top - Inches(0.55)

    for ph in master.placeholders:
        t = ph.placeholder_format.type
        if t == PP_PLACEHOLDER.TITLE:
            _place(ph, title_x, title_top, title_w, title_h)
        elif t == PP_PLACEHOLDER.BODY:
            _place(ph, M, body_top, W - 2 * M, body_h)

    for layout in prs.slide_layouts:
        phs = {ph.placeholder_format.idx: ph for ph in layout.placeholders}
        if layout.name == "Title Slide":
            _place(phs[0], M, int(H * 0.30), int(W * 0.58), Inches(1.9))
            _lst_style(phs[0], size.get("cover", 40), ink, heading, bold=True, anchor="b")
            _place(phs[1], M, int(H * 0.30) + Inches(2.05), int(W * 0.58), Inches(1.3))
            _lst_style(phs[1], size.get("subtitle", 18), muted, body)
        elif layout.name == "Section Header":
            inverted = bool(s.get("section_inverted", True))
            fg, sub = (paper, c.get("section_sub", "A0AEC0")) if inverted else (ink, soft)
            if inverted:
                _background(layout, ink)
                layout._element.set("showMasterSp", "0")
            _place(phs[0], M, int(H * 0.36), W - 2 * M, Inches(1.5))
            _lst_style(phs[0], size.get("section", 36), fg, heading, bold=True, anchor="b")
            _place(phs[1], M, int(H * 0.36) + Inches(1.6), W - 2 * M, Inches(1.0))
            _lst_style(phs[1], size.get("subtitle", 18), sub, body, anchor="t")
        elif layout.name == "Two Content":
            half = (W - 2 * M - Inches(0.35)) // 2
            _place(phs[0], title_x, title_top, title_w, title_h)
            _place(phs[1], M, body_top, half, body_h)
            _place(phs[2], M + half + Inches(0.35), body_top, half, body_h)
        else:
            for idx, ph in phs.items():
                if ph.placeholder_format.type == PP_PLACEHOLDER.TITLE:
                    _place(ph, title_x, title_top, title_w, title_h)
                elif ph.placeholder_format.type in (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT):
                    _place(ph, M, body_top, W - 2 * M, body_h)

    if logo_at == "title":
        # a title beside a logo is centred on it
        for holder in [master] + [l for l in prs.slide_layouts if l.name not in ("Title Slide", "Section Header")]:
            for ph in holder.placeholders:
                if ph.placeholder_format.type == PP_PLACEHOLDER.TITLE:
                    ph._element.find(qn("p:txBody")).find(qn("a:bodyPr")).set("anchor", "ctr")

    master_art = s.get("master_art")
    ma_path = (theme.source.parent / master_art).resolve() if master_art else None
    if ma_path and ma_path.exists():
        # faint line art in the top right of every slide, behind everything
        ma_h = int(H * float(s.get("master_art_height", 0.34)))
        from PIL import Image
        with Image.open(ma_path) as im:
            ma_w = int(ma_h * im.size[0] / im.size[1])
        pic = _place_picture(prs, master, ma_path, W - ma_w + Inches(0.2), -Inches(0.1), height=ma_h)
        spTree = master.shapes._spTree
        spTree.remove(pic._element)
        spTree.insert(2, pic._element)
    if logo_w:
        if logo_at == "bottom-left":
            _place_picture(prs, master, logo_path, M, H - Inches(0.15) - logo_h - Inches(0.05), height=logo_h)
        elif logo_at == "title":
            _place_picture(prs, master, logo_path, M, title_top + (title_h - logo_h) // 2, height=logo_h)
            if rule:
                rule_h = int(logo_h * 1.25)
                _place_rule(prs, master, M + logo_w + Inches(0.25), title_top + (title_h - rule_h) // 2,
                            rule_w, rule_h, ink)
            # the title slide's title is mid-page, so it hides the master's logo and
            # takes its own, usually the full wordmark
            cover = prs.slide_layouts.get_by_name("Title Slide")
            cover._element.set("showMasterSp", "0")
            cl = s.get("cover_logo")
            cl_path = (theme.source.parent / cl).resolve() if cl else None
            if cl_path and cl_path.exists():
                cl_h = Inches(float(s.get("cover_logo_h_in", 0.3)))
                _place_picture(prs, cover, cl_path, M, H - Inches(0.25) - cl_h, height=cl_h)
        else:
            _place_picture(prs, master, logo_path, W - M - logo_w, title_top + Inches(0.12), height=logo_h)
    section_logo = s.get("section_logo")
    sl_path = (theme.source.parent / section_logo).resolve() if section_logo else None
    if sl_path and sl_path.exists():
        # the section layout hides the master's shapes (a dark logo would vanish on the
        # inverted background), so it carries its own light one, in the logo's corner
        sl_h = Inches(float(s.get("section_logo_h_in", 0.5)))
        _place_picture(prs, prs.slide_layouts.get_by_name("Section Header"), sl_path,
                       M, H - Inches(0.2) - sl_h, height=sl_h)
    art = s.get("cover_art")
    art_path = (theme.source.parent / art).resolve() if art else None
    if art_path and art_path.exists():
        cover = prs.slide_layouts.get_by_name("Title Slide")
        art_h = int(H * float(s.get("cover_art_height", 0.85)))
        from PIL import Image
        with Image.open(art_path) as im:
            art_w = int(art_h * im.size[0] / im.size[1])
        # keep the art clear of the title: it may start no further left than the
        # title box ends, so a two-line title never runs over the lines
        room = W - (M + int(W * 0.58)) - Inches(0.2) + Inches(0.4)
        if art_w > room:
            art_h = int(art_h * room / art_w)
            art_w = room
        pic = _place_picture(prs, cover, art_path, W - art_w + Inches(0.4), H - art_h, height=art_h)
        spTree = cover.shapes._spTree
        spTree.remove(pic._element)
        spTree.insert(2, pic._element)   # behind the placeholders

    cp = prs.core_properties
    cp.title = theme.name
    cp.author = cp.last_modified_by = ""
    cp.revision = 1
    import datetime as dt
    cp.created = cp.modified = dt.datetime(2026, 1, 1)
    prs.save(str(out))
    return out


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--theme", default="default", help="theme name, or a path to a theme folder")
    args = ap.parse_args(argv)
    out = build(Theme.load(args.theme))
    print(out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
