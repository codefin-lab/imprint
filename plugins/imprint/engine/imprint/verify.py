"""Assert a generated .docx actually matches its theme — run after every build."""
from __future__ import annotations

import docx
from docx.oxml.ns import qn

from .theme import Theme


def check(path, theme: Theme) -> tuple[list[str], list[str]]:
    """Return (errors, warnings). Errors mean the file is wrong; warnings are reminders."""
    d = docx.Document(str(path))
    problems: list[str] = []
    warnings: list[str] = []

    fill = theme.table.get("header_fill", "efefef").lower()
    for i, t in enumerate(d.tables, 1):
        seen = set()
        for cell in t.rows[0].cells:
            if id(cell._tc) in seen:
                continue
            seen.add(id(cell._tc))
            tcPr = cell._tc.tcPr
            shd = tcPr.find(qn("w:shd")) if tcPr is not None else None
            if shd is None or (shd.get(qn("w:fill")) or "").lower() != fill:
                problems.append(f"table {i}: header cell not shaded {fill}")
        if t._tbl.tblPr.find(qn("w:tblBorders")) is None:
            problems.append(f"table {i}: no table-level borders")
        for row in t.rows:
            for cell in row.cells:
                tcPr = cell._tc.tcPr
                if tcPr is not None and tcPr.find(qn("w:tcBorders")) is not None:
                    problems.append(f"table {i}: a cell overrides the table border")
                    break

    half = str(theme.body_pt * 2)
    for p in d.paragraphs:
        pPr = p._p.pPr
        if pPr is None or pPr.find(qn("w:numPr")) is None:
            continue
        rPr = pPr.find(qn("w:rPr"))
        sz = rPr.find(qn("w:sz")) if rPr is not None else None
        if sz is None or sz.get(qn("w:val")) != half:
            problems.append("a list marker is not at the body size")
            break

    unnamed = [s.get(qn("w:styleId")) for s in d.styles.element.findall(qn("w:style"))
               if s.find(qn("w:name")) is None]
    if unnamed:
        problems.append(f"styles without w:name (Word will offer to repair): {unnamed}")

    leftovers = set()
    for node in d.element.body.iter(qn("w:t")):
        text = node.text or ""
        if "{{" in text and "}}" in text:
            leftovers.add(text[text.find("{{"):text.find("}}") + 2])
    if leftovers:
        warnings.append(f"{len(leftovers)} placeholder(s) still unfilled: "
                        + ", ".join(sorted(leftovers)[:6])
                        + (" ..." if len(leftovers) > 6 else ""))

    return problems, warnings
