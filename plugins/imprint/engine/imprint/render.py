"""Turn parsed Markdown blocks into a .docx, following a Theme.

The theme owns every visual decision; nothing here hard-codes a font, size or
colour.  Swap themes/<name>/theme.yaml + base.docx and the same Markdown comes
out in the new design.
"""
from __future__ import annotations

import re

import copy
import hashlib
import tempfile
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .gantt import parse_spec, render_png
from .markdown import inline_spans, parse, split_front_matter
from .theme import Theme


# --------------------------------------------------------------- xml helpers

def _el(tag: str, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn(f"w:{k}"), str(v))
    return e


def _set_fonts(run, name: str):
    rPr = run._r.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.insert(0, rf)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(attr), name)


def write_spans(paragraph, text: str, theme: Theme, *, bold=False, size=None, font=None):
    size = size or theme.body_pt
    font = font or theme.body_font
    for chunk, b, i, mono in inline_spans(text):
        run = paragraph.add_run(chunk)
        run.bold = bold or b
        run.italic = i
        run.font.size = Pt(size)
        _set_fonts(run, theme.mono_font if mono else font)


def _keep_with_next(p):
    """Do not let this paragraph be the last thing on a page.

    A bold lead-in stranded at the foot of a page, or a numbered item whose
    sub-items start overleaf, reads as a mistake.
    """
    pPr = p._p.get_or_add_pPr()
    if pPr.find(qn("w:keepNext")) is None:
        pPr.append(_el("w:keepNext", val=1))
    return p


def _space(p, theme: Theme, key: str, default_pt: float = 0, *, before_pt=None,
           line_kind: str = None):
    """Put the theme's vertical gap on a paragraph.

    Nothing in base.docx supplies one — Normal is `before: 0, after: 0` — so
    every gap in the document is decided here, from `theme.yaml -> space:`.
    """
    twips = theme.space_twips(key, default_pt)
    before = None if before_pt is None else int(round(float(before_pt) * 20))
    line = theme.line_twips(line_kind) if line_kind else None
    if not twips and not before and not line:
        return p
    pPr = p._p.get_or_add_pPr()
    for old in pPr.findall(qn("w:spacing")):
        pPr.remove(old)
    attrs = {"after": twips}
    if before:
        attrs["before"] = before
    if line:
        attrs["line"] = line
        attrs["lineRule"] = "auto"
    pPr.append(_el("w:spacing", **attrs))
    return p


def _add_body_paragraph(document, lines, theme: Theme, *, hard: bool, indent: int = 0):
    """Render one Markdown paragraph.

    `lines` is one entry per source line.  In hard mode every newline the author
    typed becomes a real line break inside the same paragraph; in soft mode the
    lines join with a space the way Markdown normally does.  A line that ends
    with two spaces or a backslash always breaks, in either mode.
    """
    p = document.add_paragraph()
    if indent:
        step = theme.list.get("indent_twips", 720)
        p._p.get_or_add_pPr().append(_el("w:ind", left=step * indent))
    for idx, (text, forced) in enumerate(lines):
        if idx:
            if hard or lines[idx - 1][1]:
                p.add_run().add_break(WD_BREAK.LINE)
            else:
                p.add_run(" ")
        write_spans(p, text, theme)
    # A one-line, wholly bold paragraph is a lead-in for what follows — the
    # source documents use them as sub-headings without making them headings.
    # Give it room above so it groups with its text instead of floating between.
    lead = (len(lines) == 1 and lines[0][0].startswith("**")
            and lines[0][0].endswith("**") and lines[0][0].count("**") == 2)
    if lead and theme.page.get("keep_lead_with_next", True):
        _keep_with_next(p)
    return _space(p, theme, "paragraph_after", 6,
                  before_pt=theme.space.get("bold_lead_before", 8) if lead else None)


# ------------------------------------------------------------------ elements

def _fresh_num_id(document, abstract_id: str, *, start: int = 1, level: int = 0) -> int:
    """Give this list a numbering definition of its very own.

    Sharing one abstractNum between several <w:num> makes Word carry the count
    forward — list two comes out 5, 6, 7, 8 instead of restarting.  A
    <w:startOverride> alone does not reliably stop it.  Cloning the whole
    abstractNum per list does, because then no two lists share any state.

    `start` is where the list begins: an ordered list the author wrote as "3."
    after an interrupting table carries on at 3 instead of dropping back to 1.
    """
    numbering = document.part.numbering_part.element

    source = None
    for a in numbering.findall(qn("w:abstractNum")):
        if a.get(qn("w:abstractNumId")) == str(abstract_id):
            source = a
            break
    if source is None:                       # theme id missing — fall back
        return int(abstract_id)

    used_abstract = [int(a.get(qn("w:abstractNumId")))
                     for a in numbering.findall(qn("w:abstractNum"))]
    new_abstract = max(used_abstract) + 1
    clone = copy.deepcopy(source)
    clone.set(qn("w:abstractNumId"), str(new_abstract))
    for tag in ("w:nsid", "w:tmpl"):         # ids Word expects to be unique
        for el in clone.findall(qn(tag)):
            clone.remove(el)
    for lvl in clone.findall(qn("w:lvl")):
        if int(lvl.get(qn("w:ilvl"))) == level:
            st = lvl.find(qn("w:start"))
            if st is None:
                st = _el("w:start", val=start)
                lvl.insert(0, st)
            else:
                st.set(qn("w:val"), str(start))

    # abstractNum elements must all precede the num elements
    nums = numbering.findall(qn("w:num"))
    if nums:
        nums[0].addprevious(clone)
    else:
        numbering.append(clone)

    used = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
    nid = (max(used) if used else 0) + 1
    num = _el("w:num", numId=nid)
    num.append(_el("w:abstractNumId", val=new_abstract))
    numbering.append(num)
    return nid


def _num_levels(document, abstract_id: str) -> int:
    """How many <w:lvl> the abstractNum defines; deeper items clamp to the last."""
    numbering = document.part.numbering_part.element
    for a in numbering.findall(qn("w:abstractNum")):
        if a.get(qn("w:abstractNumId")) == str(abstract_id):
            return max(1, len(a.findall(qn("w:lvl"))))
    return 1


def _list_item(document, text: str, num_id: int, theme: Theme, level: int = 0):
    p = document.add_paragraph()
    pPr = p._p.get_or_add_pPr()

    numPr = OxmlElement("w:numPr")
    numPr.append(_el("w:ilvl", val=level))
    numPr.append(_el("w:numId", val=num_id))
    pPr.append(numPr)

    indent = theme.list.get("indent_twips", 720)
    pPr.append(_el("w:ind",
                   left=indent * (level + 1),
                   hanging=theme.list.get("hanging_twips", 360)))

    # the bullet/number glyph inherits its size from the paragraph mark, not
    # from the runs — set it explicitly or it keeps the base document's size
    mark = OxmlElement("w:rPr")
    half = theme.body_pt * 2
    mark.append(_el("w:sz", val=half))
    mark.append(_el("w:szCs", val=half))
    pPr.append(mark)

    write_spans(p, text, theme)
    return _space(p, theme, "list_item_after", 2, line_kind="list")


ROMAN = [(10, "x"), (9, "ix"), (5, "v"), (4, "iv"), (1, "i")]


def _format_number(n: int, style: str) -> str:
    """Match the per-level formats the theme's numbering.xml uses."""
    if style == "lowerLetter":
        out = ""
        while n:
            n, r = divmod(n - 1, 26)
            out = chr(ord("a") + r) + out
        return out
    if style == "lowerRoman":
        out, left = "", n
        for value, sign in ROMAN:
            while left >= value:
                out += sign
                left -= value
        return out
    return str(n)


def _literal_item(document, text: str, marker: str, theme: Theme, level: int = 0):
    """A list item whose marker is printed text rather than a Word list number."""
    p = document.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    step = theme.list.get("indent_twips", 720)
    pPr.append(_el("w:ind",
                   left=step * (level + 1),
                   hanging=theme.list.get("hanging_twips", 360)))
    tabs = OxmlElement("w:tabs")
    tabs.append(_el("w:tab", val="left", pos=step * (level + 1)))
    pPr.append(tabs)
    write_spans(p, marker, theme)
    p.add_run().add_tab()
    write_spans(p, text, theme)
    return _space(p, theme, "list_item_after", 2, line_kind="list")


def _apply_table_look(table, theme: Theme):
    t = theme.table
    tblPr = table._tbl.tblPr

    for old in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borders.append(_el(f"w:{edge}", val="single",
                           sz=t.get("border_size", 8),
                           color=t.get("border_color", "000000")))
    tblPr.append(borders)

    margin = t.get("cell_margin_twips")
    if margin:
        for old in tblPr.findall(qn("w:tblCellMar")):
            tblPr.remove(old)
        cellmar = OxmlElement("w:tblCellMar")
        for side in ("top", "left", "bottom", "right"):
            cellmar.append(_el(f"w:{side}", w=margin, type="dxa"))
        tblPr.append(cellmar)


def _shade(cell, fill: str):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:shd")):
        tcPr.remove(old)
    tcPr.append(_el("w:shd", val="clear", fill=fill))


def _column_widths(rows: list[list[str]], cols: int, theme: Theme) -> list[int] | None:
    """Split the text column between table columns by how much each one holds.

    Equal columns are what python-docx gives you, and they read badly: a
    one-word "Include" column takes the same width as a paragraph of prose.
    Weighting by content is deterministic (same Markdown, same widths) unlike
    leaving it to Word's autofit, which differs between Word and LibreOffice.
    """
    if str(theme.table.get("column_widths", "content")).lower() != "content":
        return None
    total = theme.table.get("width_twips", 10170)
    cap = int(theme.table.get("measure_cap", 90))
    weights = []
    for c in range(cols):
        longest = max((len(r[c]) for r in rows if c < len(r)), default=1)
        weights.append(max(min(longest, cap), 6))
    lo = float(theme.table.get("min_column_share", 0.07))
    hi = float(theme.table.get("max_column_share", 0.45))
    shares = [w / sum(weights) for w in weights]
    shares = [min(max(sh, lo), hi) for sh in shares]
    shares = [sh / sum(shares) for sh in shares]
    widths = [int(total * sh) for sh in shares]

    # Never squeeze a column narrower than its longest single word, or the
    # header breaks mid-word ("Includ / e").  Take the shortfall off whichever
    # column is widest at the time, so prose columns pay for it, not labels.
    per_char = int(theme.table.get("char_twips", 115))
    padding = 2 * int(theme.table.get("cell_margin_twips", 100))
    for c in range(cols):
        longest_word = max((len(w) for r in rows if c < len(r)
                            for w in r[c].split()), default=1)
        need = longest_word * per_char + padding
        if widths[c] >= need:
            continue
        donor = max(range(cols), key=lambda i: widths[i])
        take = min(need - widths[c], max(widths[donor] - need, 0))
        widths[donor] -= take
        widths[c] += take

    widths[-1] += total - sum(widths)          # no rounding drift
    return widths


def _is_banner(row: list[str], cols: int) -> tuple[bool, str]:
    """A row carrying a group label and nothing else -> (is_banner, label).

    Word writes these as one cell merged across the table ("System Design"
    sitting above its deliverables); extraction turns the merge into one filled
    cell plus a run of blanks, which renders as a row with holes in it.

    The label must be **bold** to qualify.  Guessing from "only one cell has
    text" alone would swallow real data rows — a feature whose number cell
    happens to be empty is not a group heading.
    """
    filled = [c.strip() for c in row[:cols] if c.strip()]
    if len(filled) != 1:
        return False, ""
    text = filled[0]
    if not (text.startswith("**") and text.endswith("**") and text.count("**") == 2):
        return False, ""
    return True, text


def _make_banner(table, r: int, cols: int, theme: Theme, label: str):
    """Merge the row into one full-width cell carrying `label`.

    Work on the raw w:tc elements: once gridSpan is set, python-docx's
    `row.cells` repeats the merged cell once per column it covers, so deleting
    "the cells after the first" through that view deletes the first one again.
    """
    tr = table.rows[r]._tr
    tcs = tr.findall(qn("w:tc"))
    if not tcs:
        return
    keep, rest = tcs[0], tcs[1:]
    for tc in rest:
        tr.remove(tc)

    tcPr = keep.find(qn("w:tcPr"))
    if tcPr is None:
        tcPr = OxmlElement("w:tcPr")
        keep.insert(0, tcPr)
    for old in tcPr.findall(qn("w:gridSpan")):
        tcPr.remove(old)
    tcPr.append(_el("w:gridSpan", val=cols))
    for old in tcPr.findall(qn("w:tcW")):
        tcPr.remove(old)
    tcPr.append(_el("w:tcW", w=theme.table.get("width_twips", 10170), type="dxa"))

    cell = table.rows[r].cells[0]
    cell.paragraphs[0].text = ""
    write_spans(cell.paragraphs[0], label, theme)
    fill = theme.table.get("banner_fill")
    if fill:
        _shade(cell, fill)


_CELL_BREAK = re.compile(r"<br\s*/?>", re.I)


def _write_cell(paragraph, text: str, theme: Theme, *, bold=False):
    """Cell text; `<br>` becomes a real line break inside the cell, so a long
    cell can be laid out as a title plus one item per line."""
    for idx, part in enumerate(_CELL_BREAK.split(text)):
        if idx:
            paragraph.add_run().add_break(WD_BREAK.LINE)
        write_spans(paragraph, part.strip(), theme, bold=bold)


_ALIGN = {"left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT,
          "center": WD_ALIGN_PARAGRAPH.CENTER}


def _add_table(document, rows: list[list[str]], theme: Theme,
               aligns: list[str | None] | None = None):
    """`aligns` comes from the Markdown separator row, one per column. The
    header cell follows its column, so a right-aligned figure column has a
    right-aligned heading over it; an unset column keeps the style's own."""
    cols = max(len(r) for r in rows)
    table = document.add_table(rows=len(rows), cols=cols)
    _apply_table_look(table, theme)

    widths = _column_widths(rows, cols, theme)
    if widths:
        table.autofit = False
        grid = table._tbl.find(qn("w:tblGrid"))
        if grid is not None:
            for col, gridCol in zip(widths, grid.findall(qn("w:gridCol"))):
                gridCol.set(qn("w:w"), str(col))
        for row in table.rows:                 # Word honours tcW, not tblGrid alone
            for width, cell in zip(widths, row.cells):
                tcPr = cell._tc.get_or_add_tcPr()
                for old in tcPr.findall(qn("w:tcW")):
                    tcPr.remove(old)
                tcPr.append(_el("w:tcW", w=width, type="dxa"))

    banners = {}
    if theme.table.get("banner_rows", True):
        for r in range(1, len(rows)):
            ok, label = _is_banner(rows[r], cols)
            if ok:
                banners[r] = label

    cant_split = theme.table.get("keep_rows_whole", True)
    for r, row in enumerate(rows):
        header = r == 0
        if cant_split:
            # a row cut in half by a page break reads as a mistake; the tallest
            # row in a real 20-page proposal is about 4in, so nothing gets stranded
            trPr = table.rows[r]._tr.get_or_add_trPr()
            if trPr.find(qn("w:cantSplit")) is None:
                trPr.append(_el("w:cantSplit", val=1))
        if header and theme.table.get("repeat_header_row", True):
            trPr = table.rows[r]._tr.get_or_add_trPr()
            trPr.append(_el("w:tblHeader", val=1))
        for c in range(cols):
            cell = table.cell(r, c)
            cell.paragraphs[0].text = ""
            _write_cell(cell.paragraphs[0], row[c] if c < len(row) else "", theme,
                        bold=header and theme.table.get("header_bold", True))
            align = _ALIGN.get((aligns or [])[c] if c < len(aligns or []) else None)
            if align is not None and r not in banners:
                cell.paragraphs[0].alignment = align
            table_line = theme.line_twips("table")
            if table_line:
                cell.paragraphs[0]._p.get_or_add_pPr().append(
                    _el("w:spacing", after=0, line=table_line, lineRule="auto"))
            if header:
                _shade(cell, theme.table.get("header_fill", "efefef"))

    for r, label in banners.items():           # after the text, before returning
        _make_banner(table, r, cols, theme, label)
    return table


def _add_code(document, text: str, theme: Theme):
    """A code block: verbatim, monospace, on a light panel.  No inline Markdown,
    so an asterisk or a backtick in an API example stays what it is, and the
    block is kept on one page."""
    p = document.add_paragraph()
    size = max(theme.body_pt - 1, 7)
    for k, line in enumerate(text.split("\n")):
        if k:
            p.add_run().add_break(WD_BREAK.LINE)
        run = p.add_run(line.replace("\t", "    "))
        run.font.size = Pt(size)
        _set_fonts(run, theme.mono_font)
    pPr = p._p.get_or_add_pPr()
    pPr.append(_el("w:keepLines", val=1))
    fill = str(theme.table.get("banner_fill", "f7f7f7")).lstrip("#")
    pPr.append(_el("w:shd", val="clear", color="auto", fill=fill))
    _space(p, theme, "paragraph_after", 6)
    return p


def _add_rule(document, theme: Theme):
    p = document.add_paragraph()
    pBdr = OxmlElement("w:pBdr")
    pBdr.append(_el("w:bottom", val="single",
                    sz=theme.rule.get("size", 6),
                    color=theme.rule.get("color", "999999")))
    p._p.get_or_add_pPr().append(pBdr)


def _sect_pr(document):
    """The document-level sectPr — the page setup every section inherits."""
    return document.element.body.find(qn("w:sectPr"))


def _oriented(sect_pr, landscape: bool, *, title_page: bool = False):
    """A copy of sectPr flipped to the requested orientation.

    `titlePg` marks a section as having its own first-page header/footer — that
    is only true for the section holding the cover.  Carrying it into later
    sections makes every new section show the cover's footer instead of the
    running one.
    """
    out = copy.deepcopy(sect_pr)
    if not title_page:
        for tp in out.findall(qn("w:titlePg")):
            out.remove(tp)
        # without this each new section restarts at "Page 1"
        for pn in out.findall(qn("w:pgNumType")):
            out.remove(pn)
    pg = out.find(qn("w:pgSz"))
    if pg is not None:
        w = int(float(pg.get(qn("w:w"))))
        h = int(float(pg.get(qn("w:h"))))
        wide, tall = max(w, h), min(w, h)
        pg.set(qn("w:w"), str(wide if landscape else tall))
        pg.set(qn("w:h"), str(tall if landscape else wide))
        pg.set(qn("w:orient"), "landscape" if landscape else "portrait")
    return out


def _close_section(document, landscape: bool, *, title_page: bool = False):
    """End the current section here, carrying `landscape` for the part above.

    A sectPr inside a paragraph describes the section that ENDS at that
    paragraph, so this is what actually flips the pages before it.

    Hang it on the paragraph already there rather than adding a fresh one: an
    added paragraph belongs to the section that FOLLOWS, and an empty paragraph
    is enough to make a page — which is why a landscape chart used to be
    trailed by a blank portrait page every time.
    """
    body = document.element.body
    last = None
    for child in body.iterchildren():
        if child.tag == qn("w:sectPr"):
            break
        last = child

    def _free(el):
        if el is None or el.tag != qn("w:p"):
            return False
        pPr = el.find(qn("w:pPr"))
        return pPr is None or pPr.find(qn("w:sectPr")) is None

    p = last if _free(last) else document.add_paragraph()._p
    p.get_or_add_pPr().append(
        _oriented(_sect_pr(document), landscape, title_page=title_page))
    return p


def _png_size(path: Path) -> tuple[int, int] | tuple[None, None]:
    """Width/height straight out of the PNG IHDR — no image library needed."""
    with open(path, "rb") as fh:
        head = fh.read(24)
    if len(head) < 24 or head[:8] != b"\x89PNG\r\n\x1a\n":
        return None, None
    return int.from_bytes(head[16:20], "big"), int.from_bytes(head[20:24], "big")


def _add_gantt(document, spec_text: str, theme: Theme, *, landscape=False):
    """Render the spec to a PNG next to a temp dir and drop it in centred.

    Returns False when the block produced no chart — usually an unfilled
    {{placeholder}} in the dates.  The caller reports it: a silent empty chart
    leaves a blank landscape page, which is worse than a loud warning.
    """
    spec = parse_spec(spec_text, theme)
    if not spec.tasks:
        return False
    # the theme feeds into the drawing, so it belongs in the cache key
    key = f"{spec_text}\x00{theme.name}\x00{landscape}\x00{sorted(theme.gantt.items())}"
    digest = hashlib.sha1(key.encode("utf-8")).hexdigest()[:12]
    out = Path(tempfile.gettempdir()) / f"docgen-gantt-{digest}.png"
    try:
        render_png(spec, theme, out, landscape=landscape)
    except ImportError:
        p = document.add_paragraph()
        write_spans(p, "[gantt chart needs matplotlib — pip install matplotlib]", theme)
        return True
    width = (theme.gantt.get("doc_width_in_landscape", 10.4) if landscape
             else theme.gantt.get("doc_width_in", 6.9))
    # A tall chart placed at full column width overflows the text area, and Word
    # crops it rather than shrinking — the milestone labels along the bottom just
    # disappear.  Fit to whichever of width/height binds first.
    max_h = (theme.gantt.get("doc_height_in_landscape", 6.4) if landscape
             else theme.gantt.get("doc_height_in", 9.6))
    px_w, px_h = _png_size(out)
    if px_w and px_h and width * px_h / px_w > max_h:
        width = max_h * px_w / px_h
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(out), width=Inches(width))
    _space(p, theme, "image_after", 6)
    return True


IMAGE_TYPES = {".png", ".jpg", ".jpeg"}


def _image_box(theme: Theme, landscape: bool) -> tuple[float, float]:
    """(max width, max height) in inches for a picture — the text area.

    Falls back to the chart's numbers: both describe the same page area.
    """
    img, gantt = theme.image, theme.gantt
    if landscape:
        return (float(img.get("doc_width_in_landscape", gantt.get("doc_width_in_landscape", 10.4))),
                float(img.get("doc_height_in_landscape", gantt.get("doc_height_in_landscape", 6.4))))
    return (float(img.get("doc_width_in", gantt.get("doc_width_in", 6.9))),
            float(img.get("doc_height_in", gantt.get("doc_height_in", 9.6))))


def _add_image(document, src: str, caption: str, theme: Theme, *, landscape=False) -> str | None:
    """Drop a picture in centred, with its caption beneath.  Returns a warning or None.

    The figure's title is the document's job — a heading above, a caption below —
    so a diagram arrives as the bare figure.  A missing or unembeddable file is
    reported loudly rather than skipped: a hole in a client document is worse
    than a warning at build time.
    """
    path = Path(src)
    problem = None
    if not path.exists():
        problem = f"image not found: {src}"
    elif path.suffix.lower() == ".svg":
        problem = f"{path.name}: Word cannot embed SVG — export the diagram as PNG"
    elif path.suffix.lower() not in IMAGE_TYPES:
        problem = f"{path.name}: unsupported image type {path.suffix}"
    if problem:
        write_spans(document.add_paragraph(), f"[{problem}]", theme)
        return problem

    width, max_h = _image_box(theme, landscape)
    px_w, px_h = _png_size(path)
    if px_w and px_h:
        # never stretch a small image past print resolution, and fit whichever
        # of width/height binds first — Word crops an oversized picture
        width = min(width, px_w / float(theme.image.get("min_dpi", 150)))
        if width * px_h / px_w > max_h:
            width = max_h * px_w / px_h
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    if not caption:
        _space(p, theme, "image_after", 6)
        return None

    _keep_with_next(p)                      # a caption never starts a page alone
    _space(p, theme, "image_caption_gap", 3)
    spec = theme.image.get("caption", {}) or {}
    cp = document.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    write_spans(cp, caption, theme, size=spec.get("size", theme.body_pt - 1))
    color = str(spec.get("color", "")).lstrip("#")
    for run in cp.runs:
        if spec.get("italic", True):
            run.italic = True
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
    _space(cp, theme, "image_after", 6)
    return None


def _add_toc(document, blocks, theme: Theme, pages: dict | None, at: int = -1):
    """Write the table of contents as real paragraphs, not a Word field.

    A `TOC` field stays empty until something refreshes it, and LibreOffice does
    NOT refresh it when converting to PDF — the exported file would ship reading
    "[TOC placeholder]".  So the entries are written out, and the page numbers
    come from a first pass over the rendered PDF (see build_docx.py).  The
    placeholder pass writes the same number of lines, so pagination does not
    move between the two passes.
    """
    cfg = theme.toc
    levels = [int(l) for l in cfg.get("levels", [1])]
    step = int(cfg.get("indent_twips", 360))
    right = int(cfg.get("width_twips", theme.table.get("width_twips", 10170)))
    leader = str(cfg.get("leader", "dot"))

    # a contents lists what comes AFTER it — never the internal cover-note page
    # or anything else the author put above the marker
    start = at if at >= 0 else -1

    for idx, (kind, payload) in enumerate(blocks):
        if kind != "h" or idx < start:
            continue
        level, text = payload
        if level not in levels:
            continue
        p = document.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        depth = levels.index(level)
        if depth:
            pPr.append(_el("w:ind", left=step * depth))
        tabs = OxmlElement("w:tabs")
        tabs.append(_el("w:tab", val="right", leader=leader, pos=right))
        pPr.append(tabs)
        write_spans(p, text, theme,
                    bold=bool(cfg.get("bold_level_1", True)) and level == levels[0])
        p.add_run().add_tab()
        page = (pages or {}).get(text)
        write_spans(p, str(page) if page else "\u2014", theme)
        _space(p, theme, "toc_entry_after", 2)


def _has_deeper_next(items, idx: int) -> bool:
    """True when the following item is nested under this one."""
    return idx + 1 < len(items) and items[idx + 1][0] > items[idx][0]


# --------------------------------------------------------------------- build

def render_blocks(document, blocks, theme: Theme, *, break_before_h1=None,
                  hard_breaks=None, toc_pages=None) -> list[str]:
    if break_before_h1 is None:
        break_before_h1 = theme.break_before_h1
    if hard_breaks is None:
        hard_breaks = theme.hard_line_breaks
    first_block = True
    landscape = False            # current page orientation
    first_section = True         # only the cover's section keeps titlePg
    chart_open = False           # a theme-driven landscape chart section is still open
    warnings: list[str] = []
    fresh_page = False           # a section break just started a page of its own
    for idx, (kind, payload) in enumerate(blocks):
        next_kind = blocks[idx + 1][0] if idx + 1 < len(blocks) else None
        wants_chart_page = (kind in ("gantt", "markwhen")
                            and theme.landscape_gantt and not landscape)
        # Close a chart's landscape section only once something else follows —
        # closing it eagerly leaves an empty portrait section, which is a blank
        # page.  Two charts in a row simply share the one landscape section.
        if chart_open and not wants_chart_page:
            _close_section(document, True)
            chart_open = False

        if kind == "orient":
            want = payload == "landscape"
            if want != landscape:
                _close_section(document, landscape, title_page=first_section)
                first_section = False
                landscape = want
                fresh_page = True
            first_block = False
            continue

        # a chart on its own landscape page, when the theme asks for it
        if wants_chart_page:
            if not chart_open:
                _close_section(document, False, title_page=first_section)
                first_section = False
                chart_open = True
            if not _add_gantt(document, payload, theme, landscape=True):
                warnings.append("a chart block produced no bars "
                                "(unfilled {{placeholder}} in the dates?)")
            first_block = False
            continue

        if kind == "h":
            level, text = payload
            p = document.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            pPr.insert(0, _el("w:pStyle", val=theme.style_for(level)))
            # every section on its own page, without a stray empty paragraph.
            # A section break already started one, so breaking again there
            # would leave the new page empty.
            starts_page = level == 1 and break_before_h1 and not first_block
            if starts_page and not fresh_page:
                pPr.insert(1, _el("w:pageBreakBefore", val=1))
            # A heading at the very top of a page does not need its space-before:
            # the top margin is already there, and the style's 20pt on top of it
            # is what makes those pages sit visibly lower than a page that opens
            # with a table.
            if starts_page or (level == 1 and fresh_page):
                for old in pPr.findall(qn("w:spacing")):
                    pPr.remove(old)
                pPr.append(_el("w:spacing",
                               before=theme.space_twips("h1_before_on_new_page", 0),
                               after=theme.space_twips("h1_after", 6)))
            write_spans(p, text, theme, font=theme.heading_font, size=theme.heading_pt(level))
            first_block = False
            fresh_page = False
            continue
        elif kind == "p":
            indent, lines = payload
            para = _add_body_paragraph(document, lines, theme, hard=hard_breaks,
                                       indent=indent)
            # "We provide the following:" must not sit alone at the foot of a
            # page with its list overleaf.  Only short paragraphs: keepNext moves
            # the whole paragraph, and dragging a long one wastes half a page.
            limit = int(theme.page.get("intro_keep_max_chars", 90))
            if (next_kind in ("ul", "ol", "table")
                    and theme.page.get("keep_intro_with_list", True)
                    and sum(len(t) for t, _ in lines) <= limit):
                _keep_with_next(para)
            fresh_page = False
        elif kind in ("ul", "ol"):
            # one numbering instance per marker kind per list, so numbers keep
            # counting across an interleaved bullet sub-list
            if theme.literal_numbers:
                counters: dict[int, int] = {}
                for idx, (level, item_kind, item, number) in enumerate(payload):
                    depth = min(level, 2)
                    if item_kind == "ol":
                        counters[depth] = number if number else counters.get(depth, 0) + 1
                        for deeper in [d for d in counters if d > depth]:
                            del counters[deeper]
                        styles = theme.list.get("number_formats",
                                                ["decimal", "lowerLetter", "lowerRoman"])
                        style = styles[min(depth, len(styles) - 1)]
                        marker = f"{_format_number(counters[depth], style)}."
                    else:
                        marker = theme.list.get("bullet_char", "-")
                    last = _literal_item(document, item, marker, theme, level=depth)
                    if _has_deeper_next(payload, idx):
                        _keep_with_next(last)
                if payload:
                    _space(last, theme, "list_after", 6)
                continue

            nids: dict[str, int] = {}
            for idx, (level, item_kind, item, number) in enumerate(payload):
                key = "bullet_abstract" if item_kind == "ul" else "number_abstract"
                depth = min(level, _num_levels(document, theme.list[key]) - 1)
                if item_kind not in nids:
                    # an ordered list that opens at something other than 1 is the
                    # author continuing a list broken by a table or a paragraph
                    nids[item_kind] = _fresh_num_id(
                        document, theme.list[key],
                        start=number if item_kind == "ol" and number else 1,
                        level=depth)
                last = _list_item(document, item, nids[item_kind], theme, level=depth)
                if _has_deeper_next(payload, idx):
                    _keep_with_next(last)
            if payload:
                _space(last, theme, "list_after", 6)
        elif kind == "table":
            rows, aligns = payload
            _add_table(document, rows, theme, aligns)
            # OOXML gives a table no spacing of its own, and two tables cannot
            # touch, so the gap after one is an empty paragraph — sized from the
            # theme rather than left at full body height
            spacer = document.add_paragraph()
            gap = theme.space_twips("table_after", 6)
            pPr = spacer._p.get_or_add_pPr()
            pPr.append(_el("w:spacing", after=0, before=0, line=gap, lineRule="exact"))
            mark = OxmlElement("w:rPr")
            mark.append(_el("w:sz", val=2))
            pPr.append(mark)
        elif kind in ("gantt", "markwhen"):
            if not _add_gantt(document, payload, theme, landscape=landscape):
                warnings.append("a chart block produced no bars "
                                "(unfilled {{placeholder}} in the dates?)")
        elif kind == "image":
            src, caption = payload
            problem = _add_image(document, src, caption, theme, landscape=landscape)
            if problem:
                warnings.append(problem)
        elif kind == "code":
            _add_code(document, payload, theme)
        elif kind == "hr":
            _add_rule(document, theme)
        elif kind == "toc":
            _add_toc(document, blocks, theme, toc_pages, at=idx)
        elif kind == "pagebreak":
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        first_block = False

    # a chart section still open at the end just becomes the last section —
    # closing it would append an empty portrait page after the chart
    if chart_open:
        landscape = True
    # the trailing sectPr governs the LAST section; it only owns the cover when
    # the document was never split, otherwise it must not claim titlePg
    if not first_section or landscape:
        body = document.element.body
        old = _sect_pr(document)
        body.replace(old, _oriented(old, landscape, title_page=first_section))
    return warnings


def set_header_kind(document, text: str, theme: Theme):
    """Retitle the running header's left side.

    base.docx carries the theme's `labels.document_kind` there ("Proposal" in
    the default theme), which is wrong the moment the document is a minute of
    meeting or a spec.  Front matter `header_left` replaces it without needing
    a second base.
    """
    kind = _label(theme, "document_kind")
    for section in document.sections:
        for hf in (section.header, section.first_page_header, section.even_page_header):
            for node in hf._element.iter(qn("w:t")):
                if node.text and kind in node.text:
                    node.text = node.text.replace(kind, text)


def set_document_kind(document, text: str, theme: Theme):
    """Rename the document kind on the cover and in the running header.

    One base serves every kind of document: front matter `document_kind:
    Business Requirement Document` turns the proposal cover into a BRD cover.
    `header_left` still wins for the header when both are given.
    """
    kind = _label(theme, "document_kind")
    parts = [document.element.body]
    for section in document.sections:
        parts += [hf._element for hf in (section.header, section.first_page_header,
                                         section.even_page_header)]
    for part in parts:
        for node in part.iter(qn("w:t")):
            if node.text and node.text.strip() == kind:
                node.text = node.text.replace(kind, text)


def substitute(document, values: dict) -> int:
    """Replace {{KEY}} in the body and in every header/footer, including the cover."""
    parts = [document.element.body]
    for section in document.sections:
        for hf in (section.header, section.footer,
                   section.first_page_header, section.first_page_footer,
                   section.even_page_header, section.even_page_footer):
            parts.append(hf._element)

    hits = 0
    for part in parts:
        for node in part.iter(qn("w:t")):
            text = node.text or ""
            if "{{" not in text:
                continue
            for key, val in values.items():
                token = "{{%s}}" % key
                if token in text:
                    text = text.replace(token, "" if val is None else str(val))
                    hits += 1
            node.text = text
    return hits


def name_all_styles(document):
    """Word refuses to open a file whose styles have no w:name without 'repairing' it."""
    for style in document.styles.element.findall(qn("w:style")):
        if style.find(qn("w:name")) is None:
            style.insert(0, _el("w:name", val=style.get(qn("w:styleId"))))


def _fill(text: str, values: dict) -> str:
    for key, val in values.items():
        text = text.replace("{{%s}}" % key, "" if val is None else str(val))
    return text


def apply_page_setup(document, theme: Theme):
    """Margins and header/footer offsets from the theme, before anything renders.

    Every section created later is cloned from this sectPr, so setting it here
    is what makes the whole document — landscape pages included — agree.
    """
    page = theme.page
    body = document.element.body
    sect_pr = body.find(qn("w:sectPr"))
    if sect_pr is None:
        return
    mar = sect_pr.find(qn("w:pgMar"))
    if mar is None:
        return
    for key, attr in (("margin_top_in", "top"), ("margin_bottom_in", "bottom"),
                      ("margin_left_in", "left"), ("margin_right_in", "right"),
                      ("header_distance_in", "header"), ("footer_distance_in", "footer")):
        if key in page:
            mar.set(qn(f"w:{attr}"), str(int(round(float(page[key]) * 1440))))


def apply_line_height(document, theme: Theme):
    """Set line spacing per kind of text, on the styles themselves.

    `body` lands on Normal, so table cells, list items and anything else
    inheriting it move together; `h1`/`h2`/`h3` land on the heading styles named
    by `heading_style`.  Kinds the renderer applies per paragraph instead
    (`list`, `table`) are read at those call sites.
    """
    targets = {"Normal": theme.line_twips("body")}
    for level, style_id in theme.heading_style.items():
        targets[style_id] = theme.line_twips(f"h{level}")

    for style in document.styles.element.findall(qn("w:style")):
        line = targets.get(style.get(qn("w:styleId")))
        if not line:
            continue
        pPr = style.find(qn("w:pPr"))
        if pPr is None:
            pPr = OxmlElement("w:pPr")
            style.append(pPr)
        sp = pPr.find(qn("w:spacing"))
        if sp is None:
            sp = OxmlElement("w:spacing")
            pPr.append(sp)
        sp.set(qn("w:line"), str(line))
        sp.set(qn("w:lineRule"), "auto")


def _toc_marker(blocks) -> int:
    return next((i for i, (k, _) in enumerate(blocks) if k == "toc"), -1)


def _toc_headings(blocks) -> list[str]:
    """Every heading the contents could list — those after the marker.

    A contents lists what follows it, which is also what keeps the template's
    "INTERNAL — delete this page" heading out of it.
    """
    at = _toc_marker(blocks)
    return [v[1] for i, (k, v) in enumerate(blocks) if k == "h" and i > at]


def _toc_entries(blocks, theme: Theme) -> list[str]:
    """The subset of those the theme's `levels` actually put in the contents."""
    levels = [int(l) for l in theme.toc.get("levels", [1])]
    at = _toc_marker(blocks)
    return [v[1] for i, (k, v) in enumerate(blocks)
            if k == "h" and i > at and v[0] in levels]


def build(md_path: Path, out_path: Path, theme: Theme, toc_pages: dict | None = None) -> dict:
    meta, md = split_front_matter(Path(md_path).read_text(encoding="utf-8"))
    document = docx.Document(str(theme.base))
    apply_page_setup(document, theme)
    apply_line_height(document, theme)
    if str(meta.pop("cover", theme.page.get("cover", True))).lower() in ("false", "no", "0"):
        drop_cover(document)
    else:
        apply_cover_style(document, theme)
    blocks = parse(md, indent_spaces=int(theme.text.get("indent_spaces", 4)))
    # A chart becomes a PNG during rendering, so its {{placeholders}} have to be
    # filled here — the document-wide pass below only ever sees the image.
    # the theme's `defaults:` (company name, address) fill whatever the
    # document's own front matter leaves out
    meta = {**theme.defaults, **meta}
    blocks = [(k, _fill(v, meta)) if k in ("gantt", "markwhen") else (k, v)
              for k, v in blocks]
    # An image path is written relative to the document, and a document lives
    # with its project — so resolve it here, where the source's location is known.
    here = Path(md_path).resolve().parent
    blocks = [(k, (str((here / Path(v[0]).expanduser()).resolve()), v[1]))
              if k == "image" else (k, v) for k, v in blocks]
    override = meta.pop("page_break_before_h1", None)
    breaks = meta.pop("line_breaks", None)
    warnings = render_blocks(document, blocks, theme, break_before_h1=override,
                             hard_breaks=None if breaks is None else str(breaks).lower() == "hard",
                             toc_pages=toc_pages)
    header_left = meta.pop("header_left", None)
    if header_left:
        set_header_kind(document, str(header_left), theme)
    document_kind = meta.pop("document_kind", None)
    if document_kind:
        set_document_kind(document, str(document_kind), theme)
    filled = substitute(document, meta)
    name_all_styles(document)
    document.save(str(out_path))
    return {"blocks": len(blocks), "tables": len(document.tables),
            "placeholders": filled, "warnings": warnings,
            "has_toc": any(k == "toc" for k, _ in blocks),
            # every heading, so the page lookup can find them; and just the
            # ones the theme puts in the contents, for an honest entry count.
            # The contents page's own heading is not listed, so it is excluded
            # here too — otherwise the page lookup goes hunting for an entry
            # that was never written.
            # every heading after the marker, so the page lookup can find them,
            # and the subset the theme actually lists, for an honest count
            "headings": _toc_headings(blocks),
            "toc_entries": _toc_entries(blocks, theme)}


# ------------------------------------------------------------------- cover

LABEL_DEFAULTS = {
    # the literal text base.docx carries on its cover; a theme overrides these
    # under `labels:` when its base was exported with different wording
    "document_kind": "Proposal",
    "company": "{{COMPANY_NAME}}",
    "team": "{{COMPANY_TEAM}}",
    "legal_prefix": "Confidential",
}


def _label(theme: Theme, key: str) -> str:
    return str(theme.labels.get(key, LABEL_DEFAULTS[key]))


def cover_roles(theme: Theme):
    """(role, matcher) pairs; the first match wins.  base.docx holds nothing but
    the cover, so matching on the text of each paragraph is unambiguous."""
    kind, company, legal = (_label(theme, k) for k in ("document_kind", "company", "legal_prefix"))
    return (
        ("eyebrow", lambda t: t == kind),
        ("title",   lambda t: t == "{{PROJECT_NAME}}"),
        ("meta",    lambda t: t.startswith(("Version:", "Last revised:", "Document Status:"))),
        ("label",   lambda t: t in ("PREPARED FOR", "PREPARED BY")),
        ("value",   lambda t: t == "{{CLIENT_LEGAL_NAME}}"),
        ("brand",   lambda t: company in t and not t.startswith(legal)),
        ("legal",   lambda t: t.startswith(legal)),
    )


def _style_run(run, spec: dict, theme: Theme):
    rPr = run._r.get_or_add_rPr()
    for tag in ("w:sz", "w:szCs", "w:color", "w:b", "w:bCs", "w:i", "w:iCs",
                "w:caps", "w:spacing", "w:rFonts"):
        for old in rPr.findall(qn(tag)):
            rPr.remove(old)
    font = {"heading": theme.heading_font, "body": theme.body_font}.get(
        spec.get("font", "body"), spec.get("font"))
    rPr.append(_el("w:rFonts", ascii=font, hAnsi=font, cs=font))
    half = int(round(float(spec["size"]) * 2))
    rPr.append(_el("w:sz", val=half))
    rPr.append(_el("w:szCs", val=half))
    rPr.append(_el("w:color", val=str(spec.get("color", "000000")).lstrip("#")))
    if spec.get("bold"):
        rPr.append(_el("w:b", val=1))
        rPr.append(_el("w:bCs", val=1))
    if spec.get("italic"):
        rPr.append(_el("w:i", val=1))
        rPr.append(_el("w:iCs", val=1))
    if spec.get("caps"):
        rPr.append(_el("w:caps", val=1))
    if spec.get("letter_spacing_pt"):
        # w:spacing inside rPr is character tracking, in twentieths of a point
        rPr.append(_el("w:spacing", val=int(round(float(spec["letter_spacing_pt"]) * 20))))


def _style_cover_paragraph(p, spec: dict, theme: Theme, *, per_run=None):
    for run in p.runs:
        chosen = spec
        if per_run:
            for needle, alt in per_run:
                if needle in run.text:
                    chosen = alt
                    break
        _style_run(run, chosen, theme)
    pPr = p._p.get_or_add_pPr()
    for old in pPr.findall(qn("w:spacing")):
        pPr.remove(old)
    attrs = {"before": int(round(float(spec.get("before_pt", 0)) * 20)),
             "after": int(round(float(spec.get("after_pt", 0)) * 20))}
    if spec.get("line_pt"):
        attrs["line"] = int(round(float(spec["line_pt"]) * 20))
        attrs["lineRule"] = "exact"
    pPr.append(_el("w:spacing", **attrs))

    for old in pPr.findall(qn("w:pBdr")):
        pPr.remove(old)
    rule = spec.get("rule")
    if rule:
        bdr = OxmlElement("w:pBdr")
        bdr.append(_el("w:bottom", val="single", sz=int(rule.get("size", 6)),
                       space=int(rule.get("space", 8)),
                       color=str(rule.get("color", "111111")).lstrip("#")))
        pPr.append(bdr)


def drop_cover(document):
    """Remove base.docx's cover page and its first-page footer.

    base.docx is a proposal cover plus the styles; a minute of meeting or a
    spec wants the styles and the running header/footer but not the cover, so
    the body is emptied and titlePg cleared rather than shipping a second base.
    """
    body = document.element.body
    for child in list(body.iterchildren()):
        if child.tag == qn("w:sectPr"):
            break
        body.remove(child)
    sect_pr = body.find(qn("w:sectPr"))
    if sect_pr is None:
        return
    # without this the first page keeps looking for the cover's own footer
    for tp in sect_pr.findall(qn("w:titlePg")):
        sect_pr.remove(tp)
    for ref in sect_pr.findall(qn("w:footerReference")) + sect_pr.findall(qn("w:headerReference")):
        if ref.get(qn("w:type")) == "first":
            sect_pr.remove(ref)


def apply_cover_style(document, theme: Theme):
    """Restyle base.docx's cover from `theme.yaml -> cover:`.

    The cover ships from a Google Docs export where nearly every line is the
    same 14pt, so nothing on it has any rank.  Assigning a size, weight and
    colour per role is what turns it into a hierarchy — and keeping those
    decisions in the theme means a later base.docx export gets them too.
    """
    cover = theme.cover
    if not cover:
        return
    for p in document.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        for role, matches in cover_roles(theme):
            if not (matches(text) and cover.get(role)):
                continue
            # the team and the company name share one paragraph; the company
            # name is the brand, the team is a caption
            per_run = None
            if role == "brand" and cover.get("brand_sub"):
                per_run = [(_label(theme, "team"), cover["brand_sub"])]
            _style_cover_paragraph(p, cover[role], theme, per_run=per_run)
            break
