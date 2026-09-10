"""Turn Markdown into a .pptx deck, following a Theme's `slides:` section.

The same Markdown the document engine reads, mapped onto slides:

    front matter        TITLE (or PROJECT_NAME), SUBTITLE, DATE, PRESENTER make
                        the title slide; `cover: false` leaves it out
    # Heading           a section divider slide; a paragraph under it is its subtitle
    ## Heading          a content slide; everything until the next heading is its body
    ### Heading         a bold lead line inside the current slide
    ---  or \\pagebreak  a new slide that keeps the current title
    <!-- toc -->        an agenda slide listing the sections that follow
    <!-- notes: ... --> speaker notes for the current slide

On a content slide, text alone fills the body placeholder, so the outline stays
editable in PowerPoint.  A table, picture or timeline shares the slide with the
text: text on the left, the visual on the right; a visual alone gets the whole
content area.  Nothing here measures rendered text, so overflow is estimated
and reported as a warning; render the PDF and look before calling a deck done.
"""
from __future__ import annotations

import datetime as dt
import hashlib
import math
import re
import tempfile
from dataclasses import replace
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu, Inches, Pt

from .gantt import parse_spec, render_png
from .markdown import inline_spans, parse, split_front_matter
from .render import _fill, _is_banner
from .theme import Theme

NOTE = re.compile(r"<!--\s*notes?:\s*(.*?)-->", re.S | re.I)
NOTE_MARK = "\x00note:"
LEFTOVER = re.compile(r"\{\{[^}]+\}\}")
ALIGN = {"left": PP_ALIGN.LEFT, "right": PP_ALIGN.RIGHT, "center": PP_ALIGN.CENTER}
THAI_MARKS = re.compile(r"[ัิ-ฺ็-๎]")
VISUAL = ("table", "image", "gantt", "markwhen", "code")


# ------------------------------------------------------------------ settings

class Deck:
    """The theme's `slides:` values, with defaults for anything left out."""

    def __init__(self, theme: Theme):
        s = theme.slides or {}
        self.theme = theme
        self.dir = theme.source.parent if theme.source else Path(".")
        self.base = self.dir / s.get("base", "base.pptx")
        font = s.get("font", {})
        self.heading_font = font.get("heading", theme.heading_font)
        self.body_font = font.get("body", theme.body_font)
        self.mono_font = font.get("mono", theme.mono_font)
        size = s.get("size", {})
        self.body_pt = float(size.get("body", 18))
        self.levels_pt = [float(v) for v in size.get("levels", [self.body_pt, self.body_pt - 2, self.body_pt - 4])]
        self.table_pt = float(size.get("table", 12))
        self.caption_pt = float(size.get("caption", 11))
        self.footer_pt = float(size.get("footer", 9))
        self.chart_pt = float(size.get("chart", 11))
        self.code_pt = float(size.get("code", 14))
        c = s.get("colors", {})
        self.ink = str(c.get("ink", "1A202C"))
        self.muted = str(c.get("muted", "4A5568"))
        self.soft = str(c.get("soft", "6B7280"))
        self.paper = str(c.get("paper", "FFFFFF"))
        self.rule = str(c.get("rule", "CBD5E0"))
        t = s.get("table", {})
        self.header_fill = str(t.get("header_fill", "EFEFEF"))
        self.banner_fill = str(t.get("banner_fill", "F7F7F7"))
        self.border = str(t.get("border_color", self.rule))
        self.margin = Inches(float(s.get("margin_in", 0.6)))
        self.gutter = Inches(float(s.get("gutter_in", 0.35)))
        self.footer_h = Inches(float(s.get("footer_h_in", 0.55)))
        self.bullet = str(s.get("bullet_char", "•"))
        self.footer = str(s.get("footer", "{{COMPANY_NAME}}"))
        self.agenda_title = str(s.get("agenda_title", "Agenda"))
        self.max_bullets = int(s.get("max_bullets", 7))
        self.text_share = float(s.get("text_share", 0.4))
        self.footer_align = str(s.get("footer_align", "left"))
        self.slide_number = bool(s.get("slide_number", True))
        self.wide_aspect = float(s.get("wide_aspect", 1.6))
        badge = s.get("footer_badge")
        self.footer_badge = (self.dir / badge).resolve() if badge else None
        self.footer_badge_h = Inches(float(s.get("footer_badge_h_in", 0.42)))
        self.footer_last_line = str(s.get("footer_last_line", "soft"))
        inv = s.get("footer_badge_invert")
        self.footer_badge_invert = (self.dir / inv).resolve() if inv else None
        self.section_inverted = bool(s.get("section_inverted", True))
        self.section_sub = str(c.get("section_sub", "A0AEC0"))


def _rgb(hex_: str) -> RGBColor:
    return RGBColor.from_string(hex_.lstrip("#").upper())


# ------------------------------------------------------------------ text

def _visible_len(text: str) -> int:
    """Characters that take width: Thai vowel and tone marks sit on the base letter."""
    return len(THAI_MARKS.sub("", text))


def _font(run, name: str, size_pt: float, color: str, *, bold=False, italic=False):
    f = run.font
    f.size = Pt(size_pt)
    f.bold = bold
    f.italic = italic
    f.color.rgb = _rgb(color)
    f.name = name
    # Thai is a complex script to PowerPoint; without a:cs it falls back to the
    # theme's complex-script font instead of the one chosen here
    rPr = run._r.get_or_add_rPr()
    for old in rPr.findall(qn("a:cs")):
        rPr.remove(old)
    cs = OxmlElement("a:cs")
    cs.set("typeface", name)
    latin = rPr.find(qn("a:latin"))
    (latin.addnext(cs) if latin is not None else rPr.append(cs))


def _runs(p, text: str, deck: Deck, size_pt: float, color: str, *, bold=False, font=None):
    for chunk, b, i, mono in inline_spans(text):
        run = p.add_run()
        run.text = chunk
        _font(run, deck.mono_font if mono else (font or deck.body_font), size_pt, color,
              bold=bold or b, italic=i)


def _para_format(p, *, level=0, bullet=None, number=None, space_after_pt=6, align=None,
                 bullet_font=None, bullet_color=None):
    """Bullet, indent and spacing written on the paragraph itself, so a text box
    and a body placeholder look the same whatever the master says."""
    pPr = p._p.get_or_add_pPr()
    for child in list(pPr):
        pPr.remove(child)
    step = Inches(0.32)
    listed = bullet is not None or number is not None
    pPr.set("lvl", str(level))
    pPr.set("marL", str(int(step * (level + 1)) if listed else int(step * level)))
    pPr.set("indent", str(-int(step)) if listed else "0")
    if align is not None:
        p.alignment = align
    spc = OxmlElement("a:spcAft")
    pts = OxmlElement("a:spcPts")
    pts.set("val", str(int(space_after_pt * 100)))
    spc.append(pts)
    pPr.append(spc)
    if (bullet is not None or number is not None) and bullet_color:
        # the marker's own colour and font: without them PowerPoint takes both from the
        # first run, so an item that opens with `code` gets a grey monospace number
        clr = OxmlElement("a:buClr")
        srgb = OxmlElement("a:srgbClr")
        srgb.set("val", bullet_color)
        clr.append(srgb)
        pPr.append(clr)
    if (bullet is not None or number is not None) and bullet_font:
        font = OxmlElement("a:buFont")
        font.set("typeface", bullet_font)
        pPr.append(font)
    if number is not None:
        el = OxmlElement("a:buAutoNum")
        el.set("type", "arabicPeriod")
        el.set("startAt", str(number))
    elif bullet is not None:
        el = OxmlElement("a:buChar")
        el.set("char", bullet)
    else:
        el = OxmlElement("a:buNone")
    pPr.append(el)
    if (bullet is not None or number is not None) and (bullet_font or bullet_color):
        # LibreOffice draws the marker with the paragraph's default run properties
        # rather than buFont and buClr, so state them there as well
        d = OxmlElement("a:defRPr")
        if bullet_color:
            fill = OxmlElement("a:solidFill")
            srgb = OxmlElement("a:srgbClr")
            srgb.set("val", bullet_color)
            fill.append(srgb)
            d.append(fill)
        if bullet_font:
            for tag in ("a:latin", "a:cs"):
                f = OxmlElement(tag)
                f.set("typeface", bullet_font)
                d.append(f)
        pPr.append(d)


def _text_items(blocks) -> list[tuple]:
    """Flatten text blocks into (kind, text, level, number) lines."""
    items = []
    for kind, payload in blocks:
        if kind == "p":
            _, lines = payload
            text = " ".join(t.strip() for t, _ in lines if t.strip())
            if text:
                items.append(("p", text, 0, None))
        elif kind in ("ul", "ol"):
            for level, item_kind, text, number in payload:
                items.append((item_kind, text, level, number if item_kind == "ol" else None))
        elif kind == "h":
            items.append(("lead", payload[1], 0, None))
    return items


def _write_text(tf, items, deck: Deck):
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    first = True
    for kind, text, level, number in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        size = deck.levels_pt[min(level, len(deck.levels_pt) - 1)]
        if kind == "ul":
            _para_format(p, level=level, bullet=deck.bullet, bullet_font=deck.body_font, bullet_color=deck.ink)
        elif kind == "ol":
            _para_format(p, level=level, number=number or 1, bullet_font=deck.body_font, bullet_color=deck.ink)
        else:
            _para_format(p, level=0, space_after_pt=8 if kind == "lead" else 10)
        if kind in ("ul", "ol"):
            lead_span = next(iter(inline_spans(text)), None)
            if lead_span is not None and lead_span[3]:
                # LibreOffice draws a list marker in the first run's font whatever
                # buFont says, so an item that opens with `code` would get a grey
                # monospace number; an invisible run in the body font goes first
                lead = p.add_run()
                lead.text = "\u200b"
                _font(lead, deck.body_font, size, deck.ink)
        _runs(p, text, deck, size, deck.ink, bold=kind == "lead")


def _text_height(items, width_emu: int, deck: Deck) -> float:
    """Estimated height in inches; a rough average glyph is half an em wide."""
    total = 0.0
    width_in = width_emu / 914400
    for kind, text, level, _ in items:
        size = deck.levels_pt[min(level, len(deck.levels_pt) - 1)]
        avail = max(1.0, width_in - 0.32 * (level + (1 if kind in ("ul", "ol") else 0)))
        per_line = max(8, avail * 72 / (size * 0.5))
        plain = re.sub(r"[*`]", "", text)
        lines = max(1, math.ceil(_visible_len(plain) / per_line))
        total += lines * size * 1.2 / 72 + 8 / 72
    return total


# ------------------------------------------------------------------ visuals

def _picture_size(path: Path) -> tuple[int, int]:
    try:
        from PIL import Image
        with Image.open(path) as im:
            return im.size
    except Exception:  # noqa: BLE001 - any unreadable image falls back to a square
        return (1000, 1000)


def _add_picture(slide, path: Path, box, caption: str, deck: Deck):
    x, y, w, h = box
    cap_h = Inches(0.4) if caption else 0
    pw, ph = _picture_size(path)
    scale = min(w / pw, (h - cap_h) / ph)
    iw, ih = int(pw * scale), int(ph * scale)
    ix, iy = x + (w - iw) // 2, y + (h - cap_h - ih) // 2
    pic = slide.shapes.add_picture(str(path), ix, iy, iw, ih)
    if caption:
        pic._element.nvPicPr.cNvPr.set("descr", caption)
        tb = slide.shapes.add_textbox(x, iy + ih + Inches(0.08), w, Inches(0.32))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        _para_format(p, space_after_pt=0, align=PP_ALIGN.CENTER)
        for run_text in [caption]:
            run = p.add_run()
            run.text = run_text
            _font(run, deck.body_font, deck.caption_pt, deck.soft, italic=True)


def _cell_borders(cell, color: str, width_emu: int = 9525):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for old in tcPr.findall(qn(tag)):
            tcPr.remove(old)
    for i, tag in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
        ln = OxmlElement(tag)
        ln.set("w", str(width_emu))
        fill = OxmlElement("a:solidFill")
        clr = OxmlElement("a:srgbClr")
        clr.set("val", color)
        fill.append(clr)
        ln.append(fill)
        tcPr.insert(i, ln)


def _column_widths(rows, cols: int, total: int) -> list[int]:
    weights = []
    for c in range(cols):
        longest = max((_visible_len(re.sub(r"[*`]|<br\s*/?>", " ", r[c])) for r in rows if c < len(r)), default=1)
        weights.append(min(max(longest, 3), 40) + 3)
    raw = [w / sum(weights) for w in weights]
    share = [max(s, 0.08) for s in raw]
    norm = sum(share)
    widths = [int(total * s / norm) for s in share]
    widths[-1] += total - sum(widths)
    return widths


def _add_table(slide, rows, aligns, box, deck: Deck) -> float:
    """Returns the estimated height the table needs, in inches."""
    x, y, w, h = box
    cols = max(len(r) for r in rows)
    row_h = Pt(deck.table_pt * 2.1)
    shape = slide.shapes.add_table(len(rows), cols, x, y, w, row_h * len(rows))
    tbl = shape.table
    tblPr = tbl._tbl.tblPr
    tblPr.set("firstRow", "1")
    tblPr.set("bandRow", "0")
    for sid in tblPr.findall(qn("a:tableStyleId")):
        tblPr.remove(sid)
    for c, cw in enumerate(_column_widths(rows, cols, w)):
        tbl.columns[c].width = cw
    banners = {r for r in range(1, len(rows)) if _is_banner(rows[r], cols)[0]}
    aligns = aligns or []
    need = 0.0
    for r, row in enumerate(rows):
        tbl.rows[r].height = row_h
        header = r == 0
        tallest = 1
        for c in range(cols):
            cell = tbl.cell(r, c)
            text = row[c] if c < len(row) else ""
            _cell_borders(cell, deck.border)
            cell.fill.solid()
            cell.fill.fore_color.rgb = _rgb(deck.header_fill if header else
                                            deck.banner_fill if r in banners else deck.paper)
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            parts = re.split(r"<br\s*/?>", text, flags=re.I) or [""]
            for k, part in enumerate(parts):
                p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                align = ALIGN.get(aligns[c] if c < len(aligns) else None)
                _para_format(p, space_after_pt=0, align=None if r in banners else align)
                _runs(p, part.strip(), deck, deck.table_pt, deck.ink, bold=header)
            col_in = tbl.columns[c].width / 914400 - 0.16
            per_line = max(4, col_in * 72 / (deck.table_pt * 0.5))
            lines = sum(max(1, math.ceil(_visible_len(re.sub(r"[*`]", "", t)) / per_line)) for t in parts)
            tallest = max(tallest, lines)
        need += max(row_h / 914400, tallest * deck.table_pt * 1.2 / 72 + 0.1)
        if r in banners:
            tbl.cell(r, 0).merge(tbl.cell(r, cols - 1))
    return need


def _add_code(slide, text: str, box, deck: Deck) -> float:
    """A code block: monospace on a light panel, one paragraph per line so
    indentation survives.  Returns the estimated height it needs, in inches."""
    x, y, w, h = box
    lines = text.split("\n")
    size = deck.code_pt
    need = len(lines) * size * 1.25 / 72 + 0.3
    shape = slide.shapes.add_textbox(x, y, w, min(h, int(need * 914400)))
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(deck.banner_fill)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.15)
    tf.margin_top = tf.margin_bottom = Inches(0.1)
    for k, line in enumerate(lines):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        _para_format(p, space_after_pt=0)
        run = p.add_run()
        run.text = line.replace("\t", "    ") or " "
        _font(run, deck.mono_font, size, deck.ink)
    return need


def _is_wide(visuals, md_dir: Path, deck: Deck) -> bool:
    if len(visuals) != 1:
        return False
    kind, payload = visuals[0]
    if kind in ("gantt", "markwhen", "code"):
        return True
    if kind == "table":
        return max(len(r) for r in payload[0]) >= 4
    path = (md_dir / Path(payload[0]).expanduser()).resolve()
    if not path.exists():
        return False
    pw, ph = _picture_size(path)
    return pw / ph >= deck.wide_aspect


def _chart_png(spec_text: str, theme: Theme, deck: Deck, width_emu: int, workdir: Path) -> Path:
    width_in = width_emu / 914400
    t = replace(theme, gantt={**theme.gantt, "width_in": width_in, "font_pt": deck.chart_pt,
                              "height_per_task_in": 0.42, "dpi": 200, "landscape": False})
    # hash() is salted per process; a file name that changes would change the slide XML
    out = workdir / f"chart-{hashlib.sha1(spec_text.encode()).hexdigest()[:10]}.png"
    return render_png(parse_spec(spec_text, t), t, out)


# ------------------------------------------------------------------ slides

def _chunk(blocks) -> list[dict]:
    slides, current, sections = [], None, []
    toc_at = None
    for kind, payload in blocks:
        if kind == "h" and payload[0] == 1:
            current = {"kind": "section", "title": payload[1], "blocks": [], "notes": []}
            slides.append(current)
            sections.append(payload[1])
        elif kind == "h" and payload[0] == 2:
            current = {"kind": "content", "title": payload[1], "blocks": [], "notes": []}
            slides.append(current)
        elif kind in ("hr", "pagebreak"):
            title = current["title"] if current and current["kind"] == "content" else ""
            current = {"kind": "content", "title": title, "blocks": [], "notes": []}
            slides.append(current)
        elif kind == "toc":
            toc_at = len(slides)
            current = {"kind": "agenda", "title": "", "blocks": [], "notes": []}
            slides.append(current)
        elif kind == "orient":
            continue
        elif kind == "p" and payload[1] and payload[1][0][0].startswith(NOTE_MARK):
            if current is not None:
                current["notes"].append(payload[1][0][0][len(NOTE_MARK):].strip())
        else:
            if current is None or current["kind"] == "agenda":
                current = {"kind": "content", "title": "", "blocks": [], "notes": []}
                slides.append(current)
            current["blocks"].append((kind, payload))
    if toc_at is not None:
        slides[toc_at]["sections"] = [s["title"] for s in slides[toc_at + 1:] if s["kind"] == "section"]
    return slides


def _layout(prs, name: str):
    for layout in prs.slide_layouts:
        if layout.name == name:
            return layout
    raise KeyError(f"base.pptx has no layout named {name!r}; regenerate it with tools/make_base_pptx.py")


def _drop_empty_placeholders(slide, keep: set[int]):
    for ph in list(slide.placeholders):
        if ph.placeholder_format.idx not in keep:
            ph._element.getparent().remove(ph._element)


def _set_title(slide, text: str, deck: Deck):
    ph = slide.shapes.title
    ph.text_frame.text = ""
    p = ph.text_frame.paragraphs[0]
    for chunk, b, i, mono in inline_spans(text):
        run = p.add_run()
        run.text = chunk
        # size and colour come from the layout; only the complex-script font is set
        rPr = run._r.get_or_add_rPr()
        cs = OxmlElement("a:cs")
        cs.set("typeface", deck.heading_font)
        rPr.append(cs)


def _footer(slide, prs, text: str, deck: Deck, *, inverted: bool = False):
    """The footer on a light slide, or its inverted version on a dark one."""
    W, H = prs.slide_width, prs.slide_height
    lines = [t for t in text.split("\n") if t.strip()]
    box_h = Inches(0.19) * max(1, len(lines)) + Inches(0.1)
    y = H - Inches(0.15) - box_h
    right = deck.footer_align == "right"
    x0, anchor = deck.margin, MSO_ANCHOR.BOTTOM
    badge = deck.footer_badge_invert if inverted else deck.footer_badge
    label, name = (deck.section_sub, deck.paper) if inverted else (deck.soft, deck.ink)
    if badge and badge.exists() and not right:
        # a round badge, then the footer lines beside it, centred on it and set a
        # little high, where the eye reads two small lines as centred
        bh = deck.footer_badge_h
        by = H - Inches(0.2) - bh
        pic = slide.shapes.add_picture(str(badge), deck.margin, by, height=bh)
        x0, anchor = deck.margin + pic.width + Inches(0.12), MSO_ANCHOR.MIDDLE
        y = by + (bh - box_h) // 2 - Inches(0.04)
    if lines:
        w = int(W * 0.5)
        x = W - deck.margin - w if right else x0
        tb = slide.shapes.add_textbox(x, y, w, box_h)
        tf = tb.text_frame
        tf.vertical_anchor = anchor
        for k, line in enumerate(lines):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            gap = 2 if k < len(lines) - 1 else 0
            _para_format(p, space_after_pt=gap, align=PP_ALIGN.RIGHT if right else PP_ALIGN.LEFT)
            run = p.add_run()
            run.text = line
            emphasis = len(lines) > 1 and k == len(lines) - 1 and deck.footer_last_line == "ink"
            _font(run, deck.body_font, deck.footer_pt, name if emphasis else label)
    if not deck.slide_number:
        return
    nx = deck.margin if right else W - deck.margin - Inches(1.2)
    num = slide.shapes.add_textbox(nx, H - Inches(0.15) - Inches(0.26), Inches(1.2), Inches(0.26))
    p = num.text_frame.paragraphs[0]
    _para_format(p, space_after_pt=0, align=PP_ALIGN.LEFT if right else PP_ALIGN.RIGHT)
    fld = OxmlElement("a:fld")
    fld.set("id", "{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}")
    fld.set("type", "slidenum")
    rPr = OxmlElement("a:rPr")
    rPr.set("lang", "en-US")
    rPr.set("sz", str(int(deck.footer_pt * 100)))
    fill = OxmlElement("a:solidFill")
    clr = OxmlElement("a:srgbClr")
    clr.set("val", deck.soft)
    fill.append(clr)
    rPr.append(fill)
    for tag in ("a:latin", "a:cs"):
        el = OxmlElement(tag)
        el.set("typeface", deck.body_font)
        rPr.append(el)
    t = OxmlElement("a:t")
    t.text = str(len(prs.slides))
    fld.append(rPr)
    fld.append(t)
    p._p.append(fld)


def _content_area(slide, prs, deck: Deck):
    """The box under the title and above the footer."""
    title = slide.shapes.title
    top = (title.top + title.height + Inches(0.2)) if title is not None else deck.margin
    W, H = prs.slide_width, prs.slide_height
    return deck.margin, top, W - 2 * deck.margin, H - top - deck.footer_h


def _build_content(slide_spec, prs, deck: Deck, theme: Theme, md_dir: Path, workdir: Path,
                   warnings: list[str], n: int, footer: str):
    blocks = slide_spec["blocks"]
    text_blocks = [b for b in blocks if b[0] in ("p", "ul", "ol", "h")]
    visuals = [b for b in blocks if b[0] in VISUAL]
    items = _text_items(text_blocks)
    bullets = sum(1 for k, *_ in items if k in ("ul", "ol"))
    if bullets > deck.max_bullets:
        warnings.append(f"slide {n}: {bullets} bullets (theme max {deck.max_bullets}); split the slide")

    if visuals or not slide_spec["title"]:
        slide = prs.slides.add_slide(_layout(prs, "Title Only"))
    else:
        slide = prs.slides.add_slide(_layout(prs, "Title and Content"))
    if slide_spec["title"]:
        _set_title(slide, slide_spec["title"], deck)
        keep = {0}
    else:
        keep = set()

    x, y, w, h = _content_area(slide, prs, deck)
    if not visuals:
        body = next((ph for ph in slide.placeholders if ph.placeholder_format.idx == 1), None)
        if body is not None:
            keep.add(1)
            body.left, body.top, body.width, body.height = x, y, w, h
            _write_text(body.text_frame, items, deck)
        elif items:
            tb = slide.shapes.add_textbox(x, y, w, h)
            _write_text(tb.text_frame, items, deck)
        if _text_height(items, w, deck) > h / 914400:
            warnings.append(f"slide {n} ({slide_spec['title'] or 'untitled'}): text likely overflows; shorten or split")
    else:
        vy, vh = y, h
        if items and _is_wide(visuals, md_dir, deck):
            # a wide picture, table or timeline beside the text would shrink to
            # a strip; the text goes on top and the visual takes the full width
            th = min(max(Inches(0.6), int(_text_height(items, w, deck) * 914400)), int(h * 0.4))
            tb = slide.shapes.add_textbox(x, y, w, th)
            _write_text(tb.text_frame, items, deck)
            if _text_height(items, w, deck) > h * 0.4 / 914400:
                warnings.append(f"slide {n} ({slide_spec['title']}): text above the visual likely overflows")
            vx, vw = x, w
            vy, vh = y + th + deck.gutter // 2, h - th - deck.gutter // 2
        elif items:
            tw = int(w * deck.text_share)
            tb = slide.shapes.add_textbox(x, y, tw, h)
            _write_text(tb.text_frame, items, deck)
            if _text_height(items, tw, deck) > h / 914400:
                warnings.append(f"slide {n} ({slide_spec['title']}): text beside the visual likely overflows")
            vx, vw = x + tw + deck.gutter, w - tw - deck.gutter
        else:
            vx, vw = x, w
        y, h = vy, vh
        each = (vw - deck.gutter * (len(visuals) - 1)) // len(visuals)
        for k, (kind, payload) in enumerate(visuals):
            box = (vx + k * (each + deck.gutter), y, each, h)
            if kind == "table":
                rows, aligns = payload
                need = _add_table(slide, rows, aligns, box, deck)
                if need > h / 914400:
                    warnings.append(f"slide {n} ({slide_spec['title']}): table needs about {need:.1f} in, "
                                    f"the slide has {h / 914400:.1f}; split it")
            elif kind == "code":
                need = _add_code(slide, payload, box, deck)
                if need > box[3] / 914400:
                    warnings.append(f"slide {n} ({slide_spec['title']}): code needs about {need:.1f} in; "
                                    f"shorten it or split the slide")
            elif kind == "image":
                src, caption = payload
                path = (md_dir / Path(src).expanduser()).resolve()
                if path.exists():
                    _add_picture(slide, path, box, caption, deck)
                else:
                    warnings.append(f"slide {n}: image not found: {src}")
            else:
                png = _chart_png(payload, theme, deck, box[2], workdir)
                _add_picture(slide, png, box, "", deck)
    _drop_empty_placeholders(slide, keep)
    _footer(slide, prs, footer, deck)
    return slide


def _clear_slides(prs):
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst):
        prs.part.drop_rel(sldId.rId)
        sldIdLst.remove(sldId)


def _doc_date(meta: dict) -> dt.datetime:
    """Core properties carry a timestamp; take it from the front matter so the
    same Markdown always gives the same file."""
    for key in ("DATE", "DD MMM YYYY"):
        raw = str(meta.get(key, "")).strip()
        for fmt in ("%Y-%m-%d", "%d %b %Y", "%d %B %Y"):
            try:
                return dt.datetime.strptime(raw, fmt)
            except ValueError:
                pass
    return dt.datetime(2026, 1, 1)


def build(md_path: Path, out_path: Path, theme: Theme) -> dict:
    deck = Deck(theme)
    if not deck.base.exists():
        raise FileNotFoundError(f"theme '{theme.name}' has no slide base at {deck.base}; "
                                f"run tools/make_base_pptx.py --theme {theme.source.parent}")
    meta, md = split_front_matter(Path(md_path).read_text(encoding="utf-8"))
    meta = {**theme.defaults, **{str(k): v for k, v in meta.items()}}
    values = {k: str(v) for k, v in meta.items() if v is not None}
    md = NOTE.sub(lambda m: f"\n\n{NOTE_MARK} {' '.join(m.group(1).split())}\n\n", md)
    md = _fill(md, values)
    blocks = parse(md, indent_spaces=int(theme.text.get("indent_spaces", 4)))
    specs = _chunk(blocks)

    prs = Presentation(str(deck.base))
    _clear_slides(prs)
    warnings: list[str] = []
    md_dir = Path(md_path).resolve().parent
    footer = _fill(deck.footer, values)

    with tempfile.TemporaryDirectory() as tmp:
        workdir = Path(tmp)
        if str(meta.get("cover", True)).lower() != "false":
            slide = prs.slides.add_slide(_layout(prs, "Title Slide"))
            _set_title(slide, values.get("TITLE") or values.get("PROJECT_NAME", ""), deck)
            sub = next((ph for ph in slide.placeholders if ph.placeholder_format.idx == 1), None)
            lines = [values[k] for k in ("SUBTITLE", "PRESENTER", "DATE") if values.get(k)]
            if sub is not None and lines:
                sub.text_frame.text = "\n".join(lines)
                _drop_empty_placeholders(slide, {0, 1})
            else:
                _drop_empty_placeholders(slide, {0})
        for spec in specs:
            n = len(prs.slides) + 1
            if spec["kind"] == "section":
                slide = prs.slides.add_slide(_layout(prs, "Section Header"))
                _set_title(slide, spec["title"], deck)
                first = next((b for b in spec["blocks"] if b[0] == "p"), None)
                sub = next((ph for ph in slide.placeholders if ph.placeholder_format.idx == 1), None)
                if first is not None and sub is not None:
                    sub.text_frame.text = " ".join(t.strip() for t, _ in first[1][1])
                    _drop_empty_placeholders(slide, {0, 1})
                else:
                    _drop_empty_placeholders(slide, {0})
                if deck.section_inverted and deck.footer_badge_invert:
                    _footer(slide, prs, footer, deck, inverted=True)
                rest = [b for b in spec["blocks"] if b is not first]
                if rest:
                    warnings.append(f"slide {n} ({spec['title']}): only the first paragraph under "
                                    f"a section heading is shown; start a ## slide for the rest")
            elif spec["kind"] == "agenda":
                items = [("ol", s, 0, i) for i, s in enumerate(spec.get("sections", []), 1)]
                slide = _build_content({"title": deck.agenda_title, "blocks": []}, prs, deck, theme,
                                       md_dir, workdir, warnings, n, footer)
                body = slide.shapes.add_textbox(*_content_area(slide, prs, deck))
                _write_text(body.text_frame, items, deck)
            else:
                slide = _build_content(spec, prs, deck, theme, md_dir, workdir, warnings, n, footer)
            if spec["notes"]:
                slide.notes_slide.notes_text_frame.text = "\n".join(spec["notes"])

        leftovers = set()
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    leftovers.update(LEFTOVER.findall(shape.text_frame.text))
        if leftovers:
            warnings.append(f"{len(leftovers)} placeholder(s) still unfilled: "
                            + ", ".join(sorted(leftovers)[:6]))

        cp = prs.core_properties
        stamp = _doc_date(meta)
        cp.title = values.get("TITLE") or values.get("PROJECT_NAME", "")
        cp.author = cp.last_modified_by = values.get("COMPANY_NAME", "")
        cp.created = cp.modified = stamp
        cp.revision = 1
        prs.save(str(out_path))
    return {"slides": len(prs.slides), "warnings": warnings}
