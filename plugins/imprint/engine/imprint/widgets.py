"""Widgets: charts and infographics written as a YAML block.

    ```widget
    widget: column
    title: Applications per month
    items:
      Jul: 120
      Aug: 180
      Sep: 240
    highlight: Sep
    ```

Every widget reads the same standard fields (see FIELDS), so what is learnt for one
applies to all. There are two kinds:

    native charts   column, bar, stacked-column, stacked-bar, line, area, pie, doughnut,
                    waterfall: PowerPoint charts whose data the reader can edit
    infographics    kpi, progress, rings, funnel, timeline, cycle, hub, nested, waffle,
                    matrix: drawn from PowerPoint shapes, every piece editable

A widget draws inside a box, so it fits anywhere a picture fits: alone on a slide,
beside text, or several in a row. Colours and sizes come from the theme's
`slides.widgets` section; nothing here names a colour of its own except as a
fallback for a theme that has none.
"""
from __future__ import annotations

import io
import math
import re
import zipfile
from dataclasses import dataclass, field

import yaml
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu, Inches, Pt

EMU_IN = 914400

# ------------------------------------------------------------------ the standard

FIELDS = {
    "widget": "the widget's name (required)",
    "title": "a short heading above the widget",
    "note": "one line under the widget: the source, the period, a footnote",
    "items": "the data: a list of {label, value, ...}, or a `label: value` mapping",
    "categories": "a multi-series chart's labels along the axis",
    "series": "a multi-series chart's data: a list of {name, values}, or a `name: [values]` mapping",
    "prefix": "text before every number, such as THB or $",
    "suffix": "text after every number, such as % or M",
    "decimals": "digits after the decimal point (default: as written)",
    "max": "the value a full bar, ring or scale stands for (default 100 for percentages)",
    "highlight": "a label, or a list of labels, to pick out; everything else is muted",
    "colors": "one (one colour for every item) or each (the theme's palette in order)",
    "labels": "true or false: show the value on each bar, slice or point",
    "legend": "auto, bottom, right or none",
    "columns": "how many items sit side by side",
    "center": "the text in the middle of a hub, cycle or rings",
    "x": "a matrix's horizontal axis: [low end, high end], or its name",
    "y": "a matrix's vertical axis: [low end, high end], or its name",
    "quadrants": "a matrix's four quadrant names: top left, top right, bottom left, bottom right",
}

ITEM_FIELDS = {
    "label": "what the item is",
    "value": "its number (or, for kpi, any short text)",
    "note": "one short line about it",
    "delta": "a change beside a kpi, such as +12% or -3",
    "max": "this item's own full scale",
    "when": "a date or period, on a timeline",
    "status": "done, now or next, on a timeline",
    "total": "true for a waterfall bar that shows a running total",
    "x": "a matrix item's position left to right, 0 to 100",
    "y": "a matrix item's position bottom to top, 0 to 100",
}


class WidgetError(ValueError):
    pass


@dataclass
class Spec:
    name: str
    raw: dict
    items: list = field(default_factory=list)
    categories: list = field(default_factory=list)
    series: list = field(default_factory=list)     # [(name, [values])]

    def get(self, key, default=None):
        return self.raw.get(key, default)

    @property
    def highlight(self) -> set:
        h = self.raw.get("highlight")
        if h is None:
            return set()
        return {str(v) for v in (h if isinstance(h, list) else [h])}


def _items(raw) -> list[dict]:
    if raw is None:
        return []
    if isinstance(raw, dict):
        return [{"label": str(k), "value": v} for k, v in raw.items()]
    if not isinstance(raw, list):
        raise WidgetError("items must be a list or a `label: value` mapping")
    out = []
    for i, it in enumerate(raw, 1):
        if isinstance(it, dict):
            unknown = set(it) - set(ITEM_FIELDS)
            if unknown:
                raise WidgetError(f"item {i}: unknown field {', '.join(sorted(map(str, unknown)))}; "
                                  f"an item takes {', '.join(ITEM_FIELDS)}")
            out.append({**it, "label": str(it.get("label", ""))})
        else:
            out.append({"label": str(it), "value": None})
    return out


def _series(raw) -> list[tuple]:
    if raw is None:
        return []
    if isinstance(raw, dict):
        return [(str(k), list(v or [])) for k, v in raw.items()]
    out = []
    for i, s in enumerate(raw, 1):
        if not isinstance(s, dict) or "values" not in s:
            raise WidgetError(f"series {i} needs `name` and `values`")
        out.append((str(s.get("name", f"Series {i}")), list(s["values"])))
    return out


def parse(text: str) -> Spec:
    try:
        raw = yaml.safe_load(text)
    except yaml.YAMLError as e:
        mark = getattr(e, "problem_mark", None)
        where = f" at line {mark.line + 1}" if mark else ""
        raise WidgetError(f"the YAML does not read{where}: {getattr(e, 'problem', e)}") from None
    if not isinstance(raw, dict) or "widget" not in raw:
        raise WidgetError("a widget block starts with `widget: <name>`")
    name = str(raw["widget"]).strip().lower()
    if name not in REGISTRY:
        raise WidgetError(f"unknown widget {name!r}; choose from {', '.join(REGISTRY)}")
    w = REGISTRY[name]
    unknown = set(raw) - set(FIELDS)
    if unknown:
        raise WidgetError(f"unknown field {', '.join(sorted(map(str, unknown)))}; "
                          f"the standard fields are {', '.join(FIELDS)}")
    spec = Spec(name, raw, _items(raw.get("items")), [str(c) for c in raw.get("categories") or []],
                _series(raw.get("series")))
    if w.data == "series":
        if not spec.series and spec.items:
            spec.categories = [it["label"] for it in spec.items]
            spec.series = [(str(raw.get("title") or "Value"), [it.get("value") for it in spec.items])]
        if not spec.series:
            raise WidgetError(f"{name} needs `items` (one series) or `categories` and `series`")
        for sname, values in spec.series:
            if len(values) != len(spec.categories):
                raise WidgetError(f"series {sname!r} has {len(values)} values for "
                                  f"{len(spec.categories)} categories")
            for v in values:
                if not isinstance(v, (int, float)) or isinstance(v, bool):
                    raise WidgetError(f"series {sname!r}: {v!r} is not a number")
    else:
        if not spec.items:
            raise WidgetError(f"{name} needs `items`")
        if w.numeric:
            for it in spec.items:
                if not isinstance(it.get("value"), (int, float)) or isinstance(it.get("value"), bool):
                    raise WidgetError(f"item {it['label']!r}: value {it.get('value')!r} is not a number")
    lo, hi = w.count
    n = len(spec.categories) if w.data == "series" else len(spec.items)
    if n < lo:
        raise WidgetError(f"{name} needs at least {lo} items, got {n}")
    spec.raw.setdefault("_count_warning", f"{name} reads best with at most {hi} items, got {n}" if n > hi else "")
    return spec


# ------------------------------------------------------------------ colours and numbers

@dataclass
class Look:
    fg: str
    sub: str
    palette: list
    accent: str
    muted: str
    track: str
    grid: str
    positive: str
    negative: str
    card_fill: str
    card_border: str
    on_accent: str
    radius: float
    value_pt: float
    label_pt: float
    title_pt: float
    note_pt: float
    heading_font: str
    body_font: str


def look_for(deck, tone) -> Look:
    s = deck.theme.slides or {}
    w = dict(s.get("widgets") or {})
    if tone.dark:
        w = {**w, **(w.get("dark") or {})}
    size = s.get("size") or {}
    palette = [str(c).lstrip("#") for c in (w.get("palette") or tone.accents or s.get("accents") or [tone.accent])]
    return Look(
        fg=tone.fg, sub=tone.sub, palette=palette,
        accent=str(w.get("highlight") or palette[0]).lstrip("#"),
        muted=str(w.get("muted") or ("4A5568" if tone.dark else "CBD5E0")),
        track=str(w.get("track") or ("2D3748" if tone.dark else "EDF2F7")),
        grid=str(w.get("grid") or ("4A5568" if tone.dark else "E2E8F0")),
        positive=str(w.get("positive") or "2F855A"), negative=str(w.get("negative") or "C53030"),
        card_fill=tone.card_fill, card_border=tone.card_border, on_accent="FFFFFF", radius=tone.radius,
        value_pt=float(size.get("widget_value", 28)), label_pt=float(size.get("widget_label", 12)),
        title_pt=float(size.get("card_title", 16)), note_pt=float(size.get("caption", 11)),
        heading_font=deck.heading_font, body_font=deck.body_font)


def fmt(value, spec: Spec, *, suffix=True) -> str:
    if value is None:
        return ""
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return str(value)
    d = spec.get("decimals")
    if d is None:
        d = 0 if float(value).is_integer() else len(repr(float(value)).split(".")[1])
    text = f"{value:,.{int(d)}f}"
    return f"{spec.get('prefix', '') or ''}{text}{(spec.get('suffix', '') or '') if suffix else ''}"


def _number_format(spec: Spec, values) -> str:
    d = spec.get("decimals")
    if d is None:
        d = max((len(repr(float(v)).split(".")[1]) for v in values if not float(v).is_integer()), default=0)
    body = "#,##0" + ("." + "0" * int(d) if int(d) else "")
    pre = str(spec.get("prefix") or "")
    suf = str(spec.get("suffix") or "")
    q = lambda t: f'"{t}"' if t else ""
    return f"{q(pre)}{body}{q(suf)}"


def on(fill: str, look: "Look") -> str:
    """White or ink, whichever reads on `fill`."""
    r, g, b = (int(fill[k:k + 2], 16) / 255 for k in (0, 2, 4))
    return "FFFFFF" if 0.2126 * r + 0.7152 * g + 0.0722 * b < 0.6 else "1A202C"


def _series_colour(si: int, name: str, spec, look) -> str:
    if spec.highlight:
        return look.accent if name in spec.highlight else look.muted
    return look.palette[si % len(look.palette)]


def _colour(i: int, label: str, spec: Spec, look: Look, default: str = "each") -> str:
    if spec.highlight:
        return look.accent if label in spec.highlight else look.muted
    if (spec.get("colors") or default) == "one":
        return look.accent
    return look.palette[i % len(look.palette)]


# ------------------------------------------------------------------ drawing helpers

def _S():
    from . import slides
    return slides


def _L():
    from . import slide_layouts
    return slide_layouts


def _shape(slide, kind, x, y, w, h, fill=None, line=None, line_w=1.0):
    shp = _L()._plain(slide.shapes.add_shape(kind, int(x), int(y), int(max(w, 1)), int(max(h, 1))))
    S = _S()
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = S._rgb(fill)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = S._rgb(line)
        shp.line.width = Pt(line_w)
    else:
        shp.line.fill.background()
    return shp


def _txt(slide, x, y, w, h, paras, deck, *, anchor=MSO_ANCHOR.TOP, align=None, space=2):
    return _L()._text(slide, x, y, w, h, paras, deck, anchor=anchor, align=align, space=space)


def _h(text, pt, width) -> int:
    return _L()._height([(text, pt)], int(width))


def _line(slide, x1, y1, x2, y2, colour, width_pt=1.5, dash=False):
    from pptx.enum.shapes import MSO_CONNECTOR
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, int(x1), int(y1), int(x2), int(y2))
    _L()._plain(conn)
    conn.line.color.rgb = _S()._rgb(colour)
    conn.line.width = Pt(width_pt)
    if dash:
        ln = conn.line._get_or_add_ln()
        d = OxmlElement("a:prstDash")
        d.set("val", "dash")
        ln.append(d)
    return conn


def _ring(slide, cx, cy, r, thickness, fraction, colour, track):
    """A progress ring: a full track, then an arc from twelve o'clock, clockwise."""
    d = 2 * r
    ratio = max(0.02, min(0.5, thickness / d))
    if track:
        t = _shape(slide, MSO_SHAPE.DONUT, cx - r, cy - r, d, d, fill=track)
        t.adjustments[0] = ratio
    fraction = max(0.0, min(1.0, fraction))
    if fraction >= 0.999:
        a = _shape(slide, MSO_SHAPE.DONUT, cx - r, cy - r, d, d, fill=colour)
        a.adjustments[0] = ratio
    elif fraction > 0.001:
        a = _shape(slide, MSO_SHAPE.BLOCK_ARC, cx - r, cy - r, d, d, fill=colour)
        start = 270.0
        end = (270.0 + 360.0 * fraction) % 360.0
        # angles are stored in 60000ths of a degree; python-pptx normalises by 100000
        a.adjustments[0] = start * 0.6
        a.adjustments[1] = end * 0.6
        a.adjustments[2] = ratio


def _fill_box(box, need):
    x, y, w, h = box
    return x, y + max(0, (h - need) // 2), w, min(h, need)


# ------------------------------------------------------------------ native charts

CHART = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED, "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "stacked-column": XL_CHART_TYPE.COLUMN_STACKED, "stacked-bar": XL_CHART_TYPE.BAR_STACKED,
    "line": XL_CHART_TYPE.LINE_MARKERS, "area": XL_CHART_TYPE.AREA,
    "pie": XL_CHART_TYPE.PIE, "doughnut": XL_CHART_TYPE.DOUGHNUT,
}


def pin_workbook(chart, stamp: str = "2026-01-01T00:00:00Z"):
    """The workbook inside a chart carries the time it was written; pin it, so the same
    Markdown gives the same bytes."""
    part = chart.part.chart_workbook.xlsx_part
    if part is None:
        return
    src = zipfile.ZipFile(io.BytesIO(part.blob))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*", rb"\g<1>" + stamp.encode(), data)
            entry = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            entry.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(entry, data)
    part.blob = out.getvalue()


def _txpr(el_parent, size_pt, colour, font):
    """Text properties for a chart element: size, colour and the font, Thai included."""
    tx = OxmlElement("c:txPr")
    body = OxmlElement("a:bodyPr")
    lst = OxmlElement("a:lstStyle")
    p = OxmlElement("a:p")
    ppr = OxmlElement("a:pPr")
    d = OxmlElement("a:defRPr")
    d.set("sz", str(int(size_pt * 100)))
    fill = OxmlElement("a:solidFill")
    c = OxmlElement("a:srgbClr")
    c.set("val", colour)
    fill.append(c)
    d.append(fill)
    for tag in ("a:latin", "a:cs"):
        f = OxmlElement(tag)
        f.set("typeface", font)
        d.append(f)
    ppr.append(d)
    p.append(ppr)
    end = OxmlElement("a:endParaRPr")
    end.set("lang", "en-US")
    p.append(end)
    tx.extend([body, lst, p])
    for old in el_parent.findall(qn("c:txPr")):
        el_parent.remove(old)
    el_parent.append(tx)


def _axis_style(axis, look: Look, deck, *, visible=True, gridlines=False):
    axis.visible = visible
    axis.has_major_gridlines = gridlines
    axis.has_minor_gridlines = False
    if gridlines:
        gl = axis.major_gridlines.format.line
        gl.color.rgb = _S()._rgb(look.grid)
        gl.width = Pt(0.75)
    axis.format.line.color.rgb = _S()._rgb(look.grid)
    axis.tick_labels.font.size = Pt(deck.chart_pt)
    axis.tick_labels.font.color.rgb = _S()._rgb(look.sub)
    axis.tick_labels.font.name = look.body_font


def _chart(ctx, spec: Spec, box):
    slide, deck, look = ctx["slide"], ctx["deck"], ctx["look"]
    name = spec.name
    x, y, w, h = box
    values = [v for _, vals in spec.series for v in vals]
    cd = CategoryChartData(number_format=_number_format(spec, values))
    cd.categories = spec.categories
    for sname, vals in spec.series:
        cd.add_series(sname, vals)
    gf = slide.shapes.add_chart(CHART[name], int(x), int(y), int(w), int(h), cd)
    chart = gf.chart
    S = _S()
    chart.font.size = Pt(deck.chart_pt)
    chart.font.name = look.body_font
    chart.font.color.rgb = S._rgb(look.sub)
    chart.has_title = False
    multi = len(spec.series) > 1
    round_ = name in ("pie", "doughnut")
    legend = str(spec.get("legend", "auto")).lower()
    show_legend = legend in ("bottom", "right") or (legend == "auto" and (multi or round_))
    chart.has_legend = show_legend
    if show_legend:
        chart.legend.position = XL_LEGEND_POSITION.RIGHT if legend == "right" else XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(deck.chart_pt)
        chart.legend.font.color.rgb = S._rgb(look.sub)
        chart.legend.font.name = look.body_font
    plot = chart.plots[0]
    labels = spec.get("labels")
    if labels is None:
        labels = name not in ("line", "area") and not (name.startswith("stacked") and len(spec.categories) > 8)
    stacked = name.startswith("stacked")

    if name in ("column", "bar", "stacked-column", "stacked-bar"):
        plot.gap_width = 60 if not multi or stacked else 80
        if stacked:
            plot.overlap = 100
        elif multi:
            plot.overlap = -10
        for si, series in enumerate(plot.series):
            series.format.line.fill.background()
            if multi:
                series.format.fill.solid()
                series.format.fill.fore_color.rgb = S._rgb(_series_colour(si, spec.series[si][0], spec, look))
            else:
                for pi, cat in enumerate(spec.categories):
                    pt = series.points[pi]
                    pt.format.fill.solid()
                    colour = _colour(pi, cat, spec, look, default="one")
                    pt.format.fill.fore_color.rgb = S._rgb(colour)
        _axis_style(chart.category_axis, look, deck)
        chart.category_axis.format.line.color.rgb = S._rgb(look.grid)
        if name in ("bar", "stacked-bar"):
            chart.category_axis.reverse_order = True     # the first item on top, as it is read
        _axis_style(chart.value_axis, look, deck, visible=not labels, gridlines=not labels)
        if not stacked and all(v >= 0 for v in values):
            chart.value_axis.minimum_scale = 0
    elif name in ("line", "area"):
        for si, series in enumerate(plot.series):
            sname = spec.series[si][0]
            colour = _series_colour(si, sname, spec, look) if multi else look.accent
            if name == "line":
                series.smooth = False
                series.format.line.color.rgb = S._rgb(colour)
                series.format.line.width = Pt(2.5)
                series.marker.style = XL_MARKER_STYLE.CIRCLE
                series.marker.size = 7
                series.marker.format.fill.solid()
                series.marker.format.fill.fore_color.rgb = S._rgb(colour)
                series.marker.format.line.color.rgb = S._rgb(colour)
                if not multi and spec.highlight:
                    series.format.line.color.rgb = S._rgb(look.muted)
                    series.marker.format.fill.fore_color.rgb = S._rgb(look.muted)
                    series.marker.format.line.color.rgb = S._rgb(look.muted)
                    for pi, cat in enumerate(spec.categories):
                        if cat in spec.highlight:
                            m = series.points[pi].marker
                            m.style = XL_MARKER_STYLE.CIRCLE
                            m.size = 11
                            m.format.fill.solid()
                            m.format.fill.fore_color.rgb = S._rgb(look.accent)
                            m.format.line.color.rgb = S._rgb(look.accent)
                            dl = series.points[pi].data_label
                            dl.has_text_frame = False
                            dl.show_value = True
                            dl.position = XL_LABEL_POSITION.ABOVE
                            dl.font.size = Pt(deck.chart_pt + 1)
                            dl.font.bold = True
                            dl.font.color.rgb = S._rgb(look.accent)
            else:
                series.format.fill.solid()
                series.format.fill.fore_color.rgb = S._rgb(colour)
                series.format.line.fill.background()
        _axis_style(chart.category_axis, look, deck)
        _axis_style(chart.value_axis, look, deck, gridlines=True)
        chart.value_axis.format.line.fill.background()
    else:   # pie, doughnut
        series = plot.series[0]
        for pi, cat in enumerate(spec.categories):
            pt = series.points[pi]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = S._rgb(_colour(pi, cat, spec, look))
            pt.format.line.color.rgb = S._rgb(ctx["paper"])
            pt.format.line.width = Pt(1.5)
        if name == "doughnut":
            hole = plot._element.find(qn("c:holeSize"))
            if hole is None:
                hole = OxmlElement("c:holeSize")
                plot._element.append(hole)
            hole.set("val", "62")
        first = plot._element.find(qn("c:firstSliceAng"))
        if first is not None:
            first.set("val", "0")
    if labels:
        plot.has_data_labels = True
        dls = plot.data_labels
        dls.number_format = _number_format(spec, values)
        dls.number_format_is_linked = False
        dls.show_value = True
        inside = stacked or round_
        dls.font.size = Pt(deck.chart_pt + (1 if not inside else 0))
        dls.font.bold = not inside
        dls.font.color.rgb = S._rgb("FFFFFF" if inside else look.fg)
        dls.font.name = look.body_font
        if name in ("column", "bar"):
            dls.position = XL_LABEL_POSITION.OUTSIDE_END
        elif inside:
            dls.position = XL_LABEL_POSITION.CENTER
    pin_workbook(chart)
    return False


def _waterfall(ctx, spec: Spec, box):
    """Rises and falls on a stacked column: an invisible base lifts each bar."""
    slide, deck, look = ctx["slide"], ctx["deck"], ctx["look"]
    S = _S()
    cats, base, up, down, tot, shown = [], [], [], [], [], []
    running = 0.0
    for it in spec.items:
        v = float(it["value"])
        cats.append(it["label"])
        if it.get("total"):
            running = v if it.get("value") is not None else running
            base.append(0); up.append(0); down.append(0); tot.append(running)
            shown.append(running)
            continue
        start, end = running, running + v
        base.append(min(start, end))
        up.append(v if v >= 0 else 0)
        down.append(-v if v < 0 else 0)
        tot.append(0)
        shown.append(v)
        running = end
    x, y, w, h = box
    cd = CategoryChartData(number_format=_number_format(spec, [it["value"] for it in spec.items]))
    cd.categories = cats
    for n_, vals in (("base", base), ("up", up), ("down", down), ("total", tot)):
        cd.add_series(n_, vals)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, int(x), int(y), int(w), int(h), cd)
    chart = gf.chart
    chart.has_title = False
    chart.has_legend = False
    chart.font.size = Pt(deck.chart_pt)
    chart.font.name = look.body_font
    chart.font.color.rgb = S._rgb(look.sub)
    plot = chart.plots[0]
    plot.gap_width = 50
    plot.overlap = 100
    colours = [None, look.positive, look.negative, look.accent]
    for si, series in enumerate(plot.series):
        series.format.line.fill.background()
        if colours[si] is None:
            series.format.fill.background()
        else:
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = S._rgb(colours[si])
    _axis_style(chart.category_axis, look, deck)
    _axis_style(chart.value_axis, look, deck, visible=False)
    if spec.get("labels", True):
        for pi in range(len(cats)):
            si = 3 if tot[pi] else (1 if up[pi] else 2)
            dl = plot.series[si].points[pi].data_label
            tf = dl.text_frame
            v = shown[pi]
            sign = "" if si == 3 else ("+" if v > 0 else "")
            tf.text = sign + fmt(v, spec)
            for r in tf.paragraphs[0].runs:
                r.font.size = Pt(deck.chart_pt)
                r.font.bold = True
                r.font.color.rgb = S._rgb("FFFFFF")
                r.font.name = look.body_font
            dl.position = XL_LABEL_POSITION.CENTER
    pin_workbook(chart)
    return False


# ------------------------------------------------------------------ infographics

def _grid(n, spec, default_max=4):
    cols = int(spec.get("columns") or min(n, default_max))
    return max(1, cols), math.ceil(n / max(1, cols))


def _kpi(ctx, spec: Spec, box):
    slide, deck, look = ctx["slide"], ctx["deck"], ctx["look"]
    x, y, w, h = box
    n = len(spec.items)
    cols, rows = _grid(n, spec)
    gap = deck.gutter
    cw = (w - gap * (cols - 1)) // cols
    pad = Inches(0.25)
    vpt = look.value_pt + 8
    need_h = 0
    for it in spec.items:
        t = pad * 2 + Inches(0.1) + int(vpt * 1.25 / 72 * EMU_IN) + _h(it["label"], look.label_pt + 1, cw - 2 * pad)
        if it.get("note"):
            t += _h(it["note"], look.note_pt, cw - 2 * pad)
        need_h = max(need_h, t)
    room = (h - gap * (rows - 1)) // rows
    ch = min(room, need_h)
    top = y + max(0, (h - (ch * rows + gap * (rows - 1))) // 2)
    for i, it in enumerate(spec.items):
        cx, cy = x + (i % cols) * (cw + gap), top + (i // cols) * (ch + gap)
        colour = _colour(i, it["label"], spec, look)
        card = _L()._box(slide, cx, cy, cw, ch, fill=look.card_fill, line=look.card_border,
                         radius=look.radius, line_w=0.75)
        _shape(slide, MSO_SHAPE.RECTANGLE, cx, cy + pad, Inches(0.06), int(vpt * 1.25 / 72 * EMU_IN), fill=colour)
        vy = cy + pad
        vh = int(vpt * 1.25 / 72 * EMU_IN)
        value = fmt(it.get("value"), spec)
        delta = it.get("delta")
        dw = 0
        if delta not in (None, ""):
            dtext = str(delta) if not isinstance(delta, (int, float)) else (("+" if delta > 0 else "") + fmt(delta, spec))
            negative = dtext.strip().startswith("-")
            dcol = look.negative if negative else look.positive
            dw = Inches(0.35) + int(len(dtext) * look.label_pt * 0.68 / 72 * EMU_IN)
            dx = cx + cw - pad - dw
            dy = vy + Inches(0.05)
            tri = _shape(slide, MSO_SHAPE.ISOSCELES_TRIANGLE, dx, dy + Inches(0.07), Inches(0.14), Inches(0.12), fill=dcol)
            if negative:
                tri.rotation = 180
            _txt(slide, dx + Inches(0.2), dy, dw - Inches(0.2), Inches(0.28),
                 [(dtext, look.label_pt, dcol, True, None)], deck, space=0)
        _txt(slide, cx + pad + Inches(0.08), vy, cw - 2 * pad - dw, vh,
             [(value, vpt, look.fg, True, look.heading_font)], deck, anchor=MSO_ANCHOR.MIDDLE, space=0)
        ly = vy + vh + Inches(0.1)
        lh = _h(it["label"], look.label_pt + 1, cw - 2 * pad)
        _txt(slide, cx + pad, ly, cw - 2 * pad, lh, [(it["label"], look.label_pt + 1, look.fg, False, None)], deck, space=0)
        if it.get("note"):
            _txt(slide, cx + pad, ly + lh, cw - 2 * pad, cy + ch - pad - ly - lh,
                 [(str(it["note"]), look.note_pt, look.sub, False, None)], deck, space=0)
    return need_h * rows > h


def _scale(it, spec, default=100.0):
    m = it.get("max", spec.get("max"))
    if m is None:
        m = default if (spec.get("suffix") or "").strip() == "%" or all(
            0 <= float(i["value"]) <= 100 for i in spec.items) else max(float(i["value"]) for i in spec.items)
    return float(m) or 1.0


def _progress(ctx, spec: Spec, box):
    slide, deck, look = ctx["slide"], ctx["deck"], ctx["look"]
    x, y, w, h = box
    n = len(spec.items)
    lpt = look.label_pt + 2
    bar = Inches(0.16)
    row = min(Inches(0.95), h // n)
    need = row * n
    top = y + max(0, (h - need) // 2)
    vw = Inches(1.1)
    for i, it in enumerate(spec.items):
        ry = top + i * row
        colour = _colour(i, it["label"], spec, look, default="one")
        lh = Inches(0.32)
        _txt(slide, x, ry, w - vw, lh, [(it["label"], lpt, look.fg, False, None)], deck, space=0,
             anchor=MSO_ANCHOR.BOTTOM)
        _txt(slide, x + w - vw, ry, vw, lh, [(fmt(it["value"], spec), lpt, look.fg, True, look.heading_font)],
             deck, space=0, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.BOTTOM)
        by = ry + lh + Inches(0.08)
        frac = max(0.0, min(1.0, float(it["value"]) / _scale(it, spec)))
        _L()._box(slide, x, by, w, bar, fill=look.track, radius=0.08)
        if frac > 0:
            _L()._box(slide, x, by, max(bar, int(w * frac)), bar, fill=colour, radius=0.08)
        if it.get("note"):
            _txt(slide, x, by + bar + Inches(0.04), w, Inches(0.26),
                 [(str(it["note"]), look.note_pt, look.sub, False, None)], deck, space=0)
    return False


def _rings(ctx, spec: Spec, box):
    slide, deck, look = ctx["slide"], ctx["deck"], ctx["look"]
    x, y, w, h = box
    n = len(spec.items)
    cols, rows = _grid(n, spec, 5)
    gap = deck.gutter
    cw = (w - gap * (cols - 1)) // cols
    text_h = Inches(0.35) + (Inches(0.3) if any(it.get("note") for it in spec.items) else 0)
    room = (h - gap * (rows - 1)) // rows
    d = min(cw - Inches(0.2), room - text_h - Inches(0.15), Inches(2.4))
    cell = d + Inches(0.15) + text_h
    top = y + max(0, (h - (cell * rows + gap * (rows - 1))) // 2)
    for i, it in enumerate(spec.items):
        cx = x + (i % cols) * (cw + gap)
        cy = top + (i // cols) * (cell + gap)
        colour = _colour(i, it["label"], spec, look)
        frac = float(it["value"]) / _scale(it, spec)
        r = d // 2
        _ring(slide, cx + cw // 2, cy + r, r, max(Inches(0.1), d // 9), frac, colour, look.track)
        vpt = min(look.value_pt, d / EMU_IN * 72 / 4.2)
        _txt(slide, cx + cw // 2 - r, cy, d, d, [(fmt(it["value"], spec), vpt, look.fg, True, look.heading_font)],
             deck, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, space=0)
        paras = [(it["label"], look.label_pt + 1, look.fg, True, None)]
        if it.get("note"):
            paras.append((str(it["note"]), look.note_pt, look.sub, False, None))
        _txt(slide, cx, cy + d + Inches(0.15), cw, text_h, paras, deck, align=PP_ALIGN.CENTER, space=1)
    return False


def _funnel(ctx, spec: Spec, box):
    slide, deck, look = ctx["slide"], ctx["deck"], ctx["look"]
    x, y, w, h = box
    n = len(spec.items)
    gap = Inches(0.06)
    row = min(Inches(0.8), (h - gap * (n - 1)) // n)
    need = row * n + gap * (n - 1)
    top = y + max(0, (h - need) // 2)
    fw = int(w * 0.62)
    top_value = max(float(it["value"]) for it in spec.items) or 1.0
    notes = any(it.get("note") for it in spec.items)
    fx = x if notes else x + (w - fw) // 2
    for i, it in enumerate(spec.items):
        ry = top + i * (row + gap)
        share = max(0.28, float(it["value"]) / top_value)
        bw = int(fw * share)
        bx = fx + (fw - bw) // 2
        colour = _colour(i, it["label"], spec, look, default="one")
        _L()._box(slide, bx, ry, bw, row, fill=colour, radius=0.06)
        on_ = on(colour, look)
        _txt(slide, bx + Inches(0.1), ry, bw - Inches(0.2), row,
             [(fmt(it["value"], spec), look.label_pt + 4, on_, True, look.heading_font),
              (it["label"], look.label_pt, on_, False, None)],
             deck, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, space=0)
        if notes and it.get("note"):
            nx = fx + fw + Inches(0.3)
            _line(slide, bx + bw + Inches(0.05), ry + row // 2, nx - Inches(0.08), ry + row // 2, look.grid, 1)
            _txt(slide, nx, ry, x + w - nx, row, [(str(it["note"]), look.label_pt, look.sub, False, None)],
                 deck, anchor=MSO_ANCHOR.MIDDLE, space=0)
    return need > h


def _timeline(ctx, spec: Spec, box):
    slide, deck, look = ctx["slide"], ctx["deck"], ctx["look"]
    x, y, w, h = box
    n = len(spec.items)
    seg = w // n
    when_h = Inches(0.4)
    body_w = seg - Inches(0.15)
    body_h = max(_h(it["label"], look.label_pt + 2, body_w) + (_h(str(it.get("note", "")), look.note_pt + 1, body_w)
                                                                 if it.get("note") else 0) for it in spec.items)
    dot = Inches(0.26)
    need = when_h + dot + Inches(0.2) + body_h
    top = y + max(0, (h - need) // 2)
    ly = top + when_h + dot // 2
    _line(slide, x, ly, x + w, ly, look.grid, 2)
    now_seen = False
    for i, it in enumerate(spec.items):
        cx = x + i * seg + seg // 2
        status = str(it.get("status") or "").lower()
        colour = _colour(i, it["label"], spec, look, default="one")
        if status == "next" or (now_seen and not status):
            fill, line = ctx["paper"], look.muted
        else:
            fill, line = colour, None
        if status == "now":
            now_seen = True
            halo = _shape(slide, MSO_SHAPE.OVAL, cx - dot, ly - dot, 2 * dot, 2 * dot, fill=look.track)
        d_ = dot * (1.2 if status == "now" else 1)
        _shape(slide, MSO_SHAPE.OVAL, cx - d_ / 2, ly - d_ / 2, d_, d_, fill=fill, line=line, line_w=2)
        if it.get("when") is not None:
            _txt(slide, cx - seg // 2, top, seg, when_h - Inches(0.05),
                 [(str(it["when"]), look.label_pt + 1, colour if fill != ctx["paper"] else look.sub, True,
                   look.heading_font)], deck, anchor=MSO_ANCHOR.BOTTOM, align=PP_ALIGN.CENTER, space=0)
        paras = [(it["label"], look.label_pt + 2, look.fg, True, None)]
        if it.get("note"):
            paras.append((str(it["note"]), look.note_pt + 1, look.sub, False, None))
        _txt(slide, cx - body_w // 2, ly + dot // 2 + Inches(0.2), body_w, body_h, paras, deck,
             align=PP_ALIGN.CENTER, space=2)
    return need > h


def _around(n, cx, cy, rx, ry, start=-90.0):
    for i in range(n):
        a = math.radians(start + 360.0 * i / n)
        yield i, cx + rx * math.cos(a), cy + ry * math.sin(a), math.cos(a)


def _cycle(ctx, spec: Spec, box, hub=False):
    """cycle: stages round a ring, each numbered, with its text beside it.
    hub: a centre with its parts around it, joined by spokes."""
    slide, deck, look = ctx["slide"], ctx["deck"], ctx["look"]
    x, y, w, h = box
    n = len(spec.items)
    cx, cy = x + w // 2, y + h // 2
    R = int(min(h * 0.36, w * 0.2))
    node = int(min(Inches(0.95), R * 0.62)) if hub else Inches(0.46)
    if hub:
        for i, px, py, _ in _around(n, cx, cy, R, R):
            _line(slide, cx, cy, px, py, look.grid, 2)
    else:
        ring = _shape(slide, MSO_SHAPE.DONUT, cx - R, cy - R, 2 * R, 2 * R, fill=look.track)
        ring.adjustments[0] = 0.035
    centre = spec.get("center")
    if centre or hub:
        cd = int(R * (0.95 if hub else 1.1))
        _shape(slide, MSO_SHAPE.OVAL, cx - cd // 2, cy - cd // 2, cd, cd,
               fill=look.fg if hub else None, line=None if hub else None)
        if centre:
            _txt(slide, cx - cd // 2 + Inches(0.1), cy - cd // 2, cd - Inches(0.2), cd,
                 [(str(centre), look.label_pt + 3, ctx["paper"] if hub else look.fg, True, look.heading_font)],
                 deck, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, space=0)
    text_w = int(min(Inches(2.6), (w - 2 * R) // 2 - node))
    for i, px, py, cos in _around(n, cx, cy, R, R):
        colour = _colour(i, it_label(spec, i), spec, look)
        _shape(slide, MSO_SHAPE.OVAL, px - node // 2, py - node // 2, node, node, fill=colour,
               line=ctx["paper"], line_w=2.5)
        it = spec.items[i]
        inside = str(i + 1) if not hub else (str(it.get("value")) if it.get("value") is not None else str(i + 1))
        _txt(slide, px - node // 2, py - node // 2, node, node,
             [(inside, look.label_pt + (4 if not hub else 2), on(colour, look), True, look.heading_font)],
             deck, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, space=0)
        paras = [(it["label"], look.label_pt + 2, look.fg, True, None)]
        if it.get("note"):
            paras.append((str(it["note"]), look.note_pt + 1, look.sub, False, None))
        th = _h(it["label"], look.label_pt + 2, text_w) + (_h(str(it.get("note")), look.note_pt + 1, text_w)
                                                            if it.get("note") else 0)
        g = node // 2 + Inches(0.15)
        if abs(cos) < 0.2:          # top or bottom: text above or below the node
            below = py > cy
            ty = py + g if below else py - g - th
            _txt(slide, px - text_w // 2, ty, text_w, th, paras, deck, align=PP_ALIGN.CENTER, space=1,
                 anchor=MSO_ANCHOR.TOP if below else MSO_ANCHOR.BOTTOM)
        elif cos > 0:
            _txt(slide, px + g, py - th // 2, text_w, th, paras, deck, space=1)
        else:
            _txt(slide, px - g - text_w, py - th // 2, text_w, th, paras, deck, align=PP_ALIGN.RIGHT, space=1)
    return False


def it_label(spec, i):
    return spec.items[i]["label"]


def _nested(ctx, spec: Spec, box):
    """Nested circles, largest first (a market, then the part served, then the part won);
    the area of each follows its value, with a floor so the smallest still reads."""
    slide, deck, look = ctx["slide"], ctx["deck"], ctx["look"]
    x, y, w, h = box
    items = spec.items
    D = int(min(h * 0.94, w * 0.5))
    big = float(items[0]["value"]) or 1.0
    bottom = y + (h + D) // 2
    left = x
    rows = []
    for i, it in enumerate(items):
        d = int(D * max(0.3, math.sqrt(max(0.0, float(it["value"])) / big)))
        colour = _colour(i, it["label"], spec, look)
        _shape(slide, MSO_SHAPE.OVAL, left + (D - d) // 2, bottom - d, d, d, fill=colour, line=ctx["paper"], line_w=2)
        rows.append((it, colour, bottom - d))
    for i, (it, colour, ctop) in enumerate(rows):
        nxt = rows[i + 1][2] if i + 1 < len(rows) else bottom
        band = max(Inches(0.3), min(Inches(0.5), nxt - ctop))
        _txt(slide, left, ctop + Inches(0.05), D, band, [(fmt(it["value"], spec), look.label_pt + 3, on(colour, look), True,
             look.heading_font)], deck, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, space=0)
    lx = x + D + Inches(0.5)
    lw = w - D - Inches(0.5)
    row = min(Inches(1.1), h // len(items))
    ty = y + (h - row * len(items)) // 2
    for i, (it, colour, _) in enumerate(rows):
        ry = ty + i * row
        _shape(slide, MSO_SHAPE.OVAL, lx, ry + Inches(0.08), Inches(0.22), Inches(0.22), fill=colour)
        paras = [(f"{it['label']}  {fmt(it['value'], spec)}", look.label_pt + 2, look.fg, True, None)]
        if it.get("note"):
            paras.append((str(it["note"]), look.note_pt + 1, look.sub, False, None))
        _txt(slide, lx + Inches(0.35), ry, lw - Inches(0.35), row, paras, deck, space=2)
    return False


def _waffle(ctx, spec: Spec, box):
    slide, deck, look = ctx["slide"], ctx["deck"], ctx["look"]
    x, y, w, h = box
    n = len(spec.items)
    gap = deck.gutter
    cw = (w - gap * (n - 1)) // n
    text_h = Inches(0.9)
    side = min(cw, h - text_h, Inches(2.6))
    cell = side // 10
    top = y + max(0, (h - side - text_h) // 2)
    for i, it in enumerate(spec.items):
        colour = _colour(i, it["label"], spec, look)
        filled = round(100 * float(it["value"]) / _scale(it, spec))
        gx = x + i * (cw + gap) + (cw - cell * 10) // 2
        s = int(cell * 0.78)
        for k in range(100):
            r, c = divmod(k, 10)
            # filled from the bottom row up, left to right, the way a level rises
            _shape(slide, MSO_SHAPE.RECTANGLE, gx + c * cell, top + (9 - r) * cell, s, s,
                   fill=colour if k < filled else look.track)
        _txt(slide, x + i * (cw + gap), top + cell * 10 + Inches(0.1), cw, text_h,
             [(fmt(it["value"], spec), look.label_pt + 8, look.fg, True, look.heading_font),
              (it["label"], look.label_pt + 1, look.sub, False, None)], deck, align=PP_ALIGN.CENTER, space=0)
    return False


def _axis_names(v, default):
    if isinstance(v, list) and len(v) == 2:
        return str(v[0]), str(v[1]), ""
    return "", "", str(v or default)


def _matrix(ctx, spec: Spec, box):
    slide, deck, look = ctx["slide"], ctx["deck"], ctx["look"]
    x, y, w, h = box
    lab = Inches(0.35)
    side_w = w - lab
    gw = int(min(side_w, (h - lab) * 1.6))
    gh = h - lab
    gx = x + lab + (side_w - gw) // 2
    gy = y
    quads = list(spec.get("quadrants") or [])
    half_w, half_h = gw // 2, gh // 2
    for q in range(4):
        qx, qy = gx + (q % 2) * half_w, gy + (q // 2) * half_h
        strong = q == 1
        _shape(slide, MSO_SHAPE.RECTANGLE, qx, qy, half_w, half_h,
               fill=look.track if strong else look.card_fill, line=ctx["paper"], line_w=2)
        if q < len(quads):
            _txt(slide, qx + Inches(0.15), qy + Inches(0.1), half_w - Inches(0.3), Inches(0.35),
                 [(str(quads[q]), look.label_pt, look.sub, True, None)], deck, space=0,
                 align=PP_ALIGN.RIGHT if q % 2 else PP_ALIGN.LEFT)
    xl, xh, xname = _axis_names(spec.get("x"), "")
    yl, yh, yname = _axis_names(spec.get("y"), "")
    _txt(slide, gx, gy + gh + Inches(0.05), gw, lab, [(xl, look.note_pt, look.sub, False, None)], deck, space=0)
    _txt(slide, gx, gy + gh + Inches(0.05), gw, lab, [(xname, look.label_pt, look.fg, True, None)], deck,
         align=PP_ALIGN.CENTER, space=0)
    _txt(slide, gx, gy + gh + Inches(0.05), gw, lab, [(xh, look.note_pt, look.sub, False, None)], deck,
         align=PP_ALIGN.RIGHT, space=0)
    ylab = _txt(slide, gx - lab - gh // 2 + lab // 2, gy + gh // 2 - lab // 2, gh, lab,
                [(" ".join(t for t in (yl, yname or "", yh) if t) if not yname else yname,
                  look.label_pt, look.fg, True, None)], deck, align=PP_ALIGN.CENTER, space=0)
    ylab.rotation = 270
    dot = Inches(0.24)
    for i, it in enumerate(spec.items):
        px = gx + int(gw * max(0.0, min(100.0, float(it.get("x", 50)))) / 100)
        py = gy + gh - int(gh * max(0.0, min(100.0, float(it.get("y", 50)))) / 100)
        colour = _colour(i, it["label"], spec, look)
        _shape(slide, MSO_SHAPE.OVAL, px - dot // 2, py - dot // 2, dot, dot, fill=colour, line=ctx["paper"], line_w=1.5)
        right = px < gx + gw * 0.75
        tw = Inches(1.8)
        _txt(slide, px + dot if right else px - dot - tw, py - Inches(0.16), tw, Inches(0.32),
             [(it["label"], look.label_pt, look.fg, True, None)], deck, anchor=MSO_ANCHOR.MIDDLE, space=0,
             align=PP_ALIGN.LEFT if right else PP_ALIGN.RIGHT)
    return False


# ------------------------------------------------------------------ the registry

@dataclass
class Widget:
    name: str
    summary: str
    draw: object
    data: str = "items"            # items, or series (items become one series)
    numeric: bool = True
    count: tuple = (1, 8)
    wide: object = False           # bool, or f(spec) -> bool
    example: str = ""


def _many(k):
    return lambda spec: len(spec.items or spec.categories) >= k


REGISTRY: dict[str, Widget] = {}


def _reg(*ws):
    for w in ws:
        REGISTRY[w.name] = w


_reg(
    Widget("column", "vertical bars; pick one out with highlight", _chart, "series", count=(2, 12), wide=_many(7),
           example="widget: column\ntitle: Applications per month\nitems:\n  Jul: 120\n  Aug: 180\n  Sep: 240\n"
                   "highlight: Sep\nnote: Branch and app channels"),
    Widget("bar", "horizontal bars, for long labels or a ranking", _chart, "series", count=(2, 10),
           example="widget: bar\ntitle: Time to open an account\nsuffix: \" days\"\nitems:\n  Branch: 5\n"
                   "  Call centre: 3\n  App: 1\nhighlight: App"),
    Widget("stacked-column", "vertical bars split into parts", _chart, "series", count=(2, 12), wide=_many(7),
           example="widget: stacked-column\ncategories: [Q1, Q2, Q3]\nseries:\n  App: [40, 70, 110]\n"
                   "  Branch: [80, 75, 60]"),
    Widget("stacked-bar", "horizontal bars split into parts", _chart, "series", count=(2, 10),
           example="widget: stacked-bar\ncategories: [North, Central, South]\nseries:\n  Approved: [70, 82, 64]\n"
                   "  Returned: [30, 18, 36]\nsuffix: \"%\""),
    Widget("line", "a trend over time; several series compare", _chart, "series", count=(3, 24), wide=True,
           example="widget: line\ncategories: [Jan, Feb, Mar, Apr, May, Jun]\nseries:\n  2025: [12, 14, 13, 15, 16, 18]\n"
                   "  2026: [14, 17, 19, 22, 26, 31]\nhighlight: \"2026\""),
    Widget("area", "a trend with its volume", _chart, "series", count=(3, 24), wide=True,
           example="widget: area\ncategories: [Jan, Feb, Mar, Apr, May, Jun]\nseries:\n  Users: [2, 3, 5, 8, 12, 17]\n"
                   "suffix: K"),
    Widget("pie", "parts of one whole, up to six", _chart, "series", count=(2, 6),
           example="widget: pie\ntitle: Where applications come from\nsuffix: \"%\"\nitems:\n  App: 58\n"
                   "  Branch: 30\n  Partners: 12"),
    Widget("doughnut", "parts of one whole, with room for a total", _chart, "series", count=(2, 6),
           example="widget: doughnut\nsuffix: \"%\"\nitems:\n  Retail: 64\n  SME: 26\n  Corporate: 10"),
    Widget("waterfall", "how a total rises and falls, step by step", _waterfall, count=(3, 12), wide=_many(6),
           example="widget: waterfall\nprefix: \"THB \"\nsuffix: M\nitems:\n  - {label: 2025, value: 120, total: true}\n"
                   "  - {label: New clients, value: 45}\n  - {label: Price change, value: -12}\n"
                   "  - {label: Lost clients, value: -8}\n  - {label: 2026, value: 145, total: true}"),
    Widget("kpi", "headline numbers with their change", _kpi, numeric=False, count=(1, 8), wide=_many(3),
           example="widget: kpi\nitems:\n  - {label: Accounts opened, value: \"12,400\", delta: +18%, note: vs 2025}\n"
                   "  - {label: Days to open, value: 1, delta: -4, note: was 5}\n"
                   "  - {label: Customer rating, value: 4.6, delta: +0.8}"),
    Widget("progress", "bars against a target or a full scale", _progress, count=(1, 7),
           example="widget: progress\nsuffix: \"%\"\nitems:\n  - {label: Requirements, value: 100}\n"
                   "  - {label: Build, value: 72}\n  - {label: Testing, value: 35}\n  - {label: Rollout, value: 0}"),
    Widget("rings", "percentages as rings, one to five", _rings, count=(1, 5), wide=_many(3),
           example="widget: rings\nsuffix: \"%\"\nitems:\n  - {label: Straight through, value: 82}\n"
                   "  - {label: Returned, value: 11}\n  - {label: Rejected, value: 7}"),
    Widget("funnel", "stages that narrow, with the number at each", _funnel, count=(2, 6),
           example="widget: funnel\nitems:\n  - {label: Visited, value: 12000}\n  - {label: Started, value: 4200}\n"
                   "  - {label: Submitted, value: 2600, note: 62% of starts}\n  - {label: Opened, value: 2100}"),
    Widget("timeline", "milestones along a line: done, now, next", _timeline, numeric=False, count=(2, 7),
           wide=True,
           example="widget: timeline\nitems:\n  - {when: Q1, label: Discovery, status: done}\n"
                   "  - {when: Q2, label: Build, status: done}\n  - {when: Q3, label: Pilot, status: now}\n"
                   "  - {when: Q4, label: Rollout, status: next}"),
    Widget("cycle", "stages that repeat, round a ring", lambda c, s, b: _cycle(c, s, b), numeric=False,
           count=(3, 8), wide=True,
           example="widget: cycle\ncenter: Strategy\nitems:\n  - {label: Scan, note: Markets and risks}\n"
                   "  - {label: Formulate, note: Goals and choices}\n  - {label: Implement, note: Plans and budget}\n"
                   "  - {label: Evaluate, note: Measure and adjust}"),
    Widget("hub", "one centre and the parts around it", lambda c, s, b: _cycle(c, s, b, hub=True), numeric=False,
           count=(3, 8), wide=True,
           example="widget: hub\ncenter: Platform\nitems:\n  - {label: Onboarding}\n  - {label: Identity}\n"
                   "  - {label: Screening}\n  - {label: Accounts}\n  - {label: Reports}"),
    Widget("nested", "a whole and the parts inside it, such as a market", _nested, count=(2, 4), wide=True,
           example="widget: nested\nprefix: \"THB \"\nsuffix: B\nitems:\n  - {label: Total market, value: 48}\n"
                   "  - {label: Served market, value: 12}\n  - {label: Target share, value: 2.4}"),
    Widget("waffle", "a share as a hundred squares", _waffle, count=(1, 4), wide=_many(3),
           example="widget: waffle\nsuffix: \"%\"\nitems:\n  - {label: Apply in the app, value: 58}\n"
                   "  - {label: Apply at a branch, value: 30}"),
    Widget("matrix", "items placed on two axes, in four quadrants", _matrix, numeric=False, count=(1, 12),
           wide=True,
           example="widget: matrix\nx: [Low effort, High effort]\ny: Value\nquadrants: [Quick wins, Big bets, "
                   "Fill-ins, Money pits]\nitems:\n  - {label: e-KYC, x: 25, y: 85}\n"
                   "  - {label: Core upgrade, x: 85, y: 80}\n  - {label: New branding, x: 30, y: 30}"),
)


# ------------------------------------------------------------------ entry points

def is_wide(text: str) -> bool:
    try:
        spec = parse(text)
    except WidgetError:
        return False
    w = REGISTRY[spec.name].wide
    return w(spec) if callable(w) else bool(w)


def render(slide, text: str, box, deck, tone, warnings: list, where: str, paper: str | None = None):
    """Draw the widget in `box`; problems become warnings naming `where`."""
    try:
        spec = parse(text)
    except WidgetError as e:
        warnings.append(f"{where}: widget: {e}")
        return
    if spec.raw.get("_count_warning"):
        warnings.append(f"{where}: {spec.raw['_count_warning']}")
    look = look_for(deck, tone)
    x, y, w, h = box
    title, note = spec.get("title"), spec.get("note")
    if title:
        th = _h(str(title), look.title_pt, w)
        _txt(slide, x, y, w, th, [(str(title), look.title_pt, look.fg, True, look.heading_font)], deck, space=0)
        y, h = y + th + Inches(0.12), h - th - Inches(0.12)
    if note:
        nh = _h(str(note), look.note_pt, w)
        _txt(slide, x, y + h - nh, w, nh, [(str(note), look.note_pt, look.sub, False, None)], deck, space=0)
        h -= nh + Inches(0.08)
    ctx = {"slide": slide, "deck": deck, "look": look, "tone": tone,
           "paper": paper or (tone.card_fill if tone.dark else deck.paper)}
    if tone.dark:
        dk = (deck.theme.slides or {}).get("dark") or {}
        ctx["paper"] = str(dk.get("background", "1A202C"))
    try:
        crowded = REGISTRY[spec.name].draw(ctx, spec, (x, y, w, h))
    except (TypeError, ValueError, KeyError) as e:
        warnings.append(f"{where}: widget {spec.name}: {e}")
        return
    if crowded:
        warnings.append(f"{where}: widget {spec.name} is crowded in its space; fewer items or a slide of its own")


def as_table(text: str):
    """(title, rows, aligns, note) for a document, which shows a widget's data as a table."""
    spec = parse(text)
    w = REGISTRY[spec.name]
    if w.data == "series":
        rows = [[""] + [n for n, _ in spec.series]]
        for ci, cat in enumerate(spec.categories):
            rows.append([cat] + [fmt(vals[ci], spec) for _, vals in spec.series])
        aligns = ["left"] + ["right"] * len(spec.series)
    else:
        cols = [k for k in ("when", "label", "value", "delta", "note") if any(it.get(k) not in (None, "") for it in spec.items)]
        head = {"when": "When", "label": "Item", "value": "Value", "delta": "Change", "note": "Note"}
        rows = [[head[c] for c in cols]]
        for it in spec.items:
            rows.append([fmt(it.get(c), spec) if c == "value" else str(it.get(c) if it.get(c) is not None else "")
                         for c in cols])
        aligns = ["right" if c in ("value", "delta") else "left" for c in cols]
    return spec.get("title"), rows, aligns, spec.get("note")


def catalogue() -> str:
    lines = []
    for w in REGISTRY.values():
        kind = "chart" if w.draw in (_chart, _waterfall) else "shapes"
        lines.append(f"{w.name:<15} {kind:<7} {w.summary}")
    return "\n".join(lines)
