"""Icons as native PowerPoint shapes.

The bundled set is Lucide (ISC licence, see icons/lucide/LICENSE): line icons on a 24 by 24
grid. Each icon becomes one freeform shape whose outline is the icon, so it stays sharp at
any size, takes any colour, and can be recoloured or resized in PowerPoint like any shape.

A theme can add its own icons: `slides.icons` names a folder of SVG files drawn the same
way (24 by 24, strokes, no fills), and a name there wins over the bundled one.
"""
from __future__ import annotations

import gzip
import json
import math
import re
import xml.etree.ElementTree as ET
from functools import lru_cache
from pathlib import Path

from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu, Pt

HERE = Path(__file__).resolve().parent / "icons" / "lucide"
GRID = 24.0
UNIT = 1000                       # path units per grid unit; DrawingML coordinates are integers
NUM = re.compile(r"[-+]?(?:\d+\.?\d*|\.\d+)(?:[eE][-+]?\d+)?")
TOKEN = re.compile(r"[MmLlHhVvCcSsQqTtAaZz]|[-+]?(?:\d+\.?\d*|\.\d+)(?:[eE][-+]?\d+)?")


@lru_cache(maxsize=1)
def bundled() -> dict:
    return json.loads(gzip.decompress((HERE / "icons.json.gz").read_bytes()))


@lru_cache(maxsize=1)
def tags() -> dict:
    return json.loads(gzip.decompress((HERE / "tags.json.gz").read_bytes()))


@lru_cache(maxsize=8)
def _folder(path: str) -> dict:
    out = {}
    for svg in sorted(Path(path).glob("*.svg")):
        root = ET.fromstring(svg.read_text(encoding="utf-8"))
        nodes = []
        for el in root.iter():
            tag = el.tag.split("}")[-1]
            if tag in ("path", "circle", "rect", "line", "ellipse", "polyline", "polygon"):
                nodes.append([tag, dict(el.attrib)])
        out[svg.stem] = nodes
    return out


def lookup(name: str, extra: str | None = None):
    name = str(name).strip().lower()
    if extra:
        found = _folder(extra).get(name)
        if found is not None:
            return found
    return bundled().get(name)


def search(words: list[str], limit: int = 60) -> list[str]:
    words = [w.lower() for w in words]
    scored = []
    for name, tg in tags().items():
        hay = [name] + tg
        score = 0
        for w in words:
            if w == name:
                score += 10
            elif w in name.split("-"):
                score += 6
            elif any(w == t for t in tg):
                score += 4
            elif w in name or any(w in t for t in tg):
                score += 1
            else:
                score = -1
                break
        if score > 0:
            scored.append((-score, name))
    found = [n for _, n in sorted(scored)[:limit]]
    if not found and len(words) > 1:
        # no icon has every word: fall back to icons with any of them
        merged = {}
        for w in words:
            for rank, n in enumerate(search([w], limit)):
                merged[n] = min(merged.get(n, limit), rank)
        found = sorted(merged, key=lambda n: (merged[n], n))[:limit]
    return found


def suggest(name: str, n: int = 5) -> list[str]:
    parts = re.split(r"[-\s]+", str(name).lower())
    return search(parts[:1], n) if parts and parts[0] else []


# ------------------------------------------------------------------ geometry

def _arc(x1, y1, rx, ry, phi, large, sweep, x2, y2):
    """An SVG arc as cubic Beziers: [(c1x, c1y, c2x, c2y, x, y)]."""
    if rx == 0 or ry == 0 or (x1 == x2 and y1 == y2):
        return [(x1, y1, x2, y2, x2, y2)]
    rx, ry = abs(rx), abs(ry)
    cp, sp = math.cos(math.radians(phi)), math.sin(math.radians(phi))
    dx, dy = (x1 - x2) / 2, (y1 - y2) / 2
    x1p, y1p = cp * dx + sp * dy, -sp * dx + cp * dy
    lam = x1p ** 2 / rx ** 2 + y1p ** 2 / ry ** 2
    if lam > 1:
        rx, ry = rx * math.sqrt(lam), ry * math.sqrt(lam)
    num = rx ** 2 * ry ** 2 - rx ** 2 * y1p ** 2 - ry ** 2 * x1p ** 2
    den = rx ** 2 * y1p ** 2 + ry ** 2 * x1p ** 2
    co = math.sqrt(max(0.0, num / den)) if den else 0.0
    if large == sweep:
        co = -co
    cxp, cyp = co * rx * y1p / ry, -co * ry * x1p / rx
    cx = cp * cxp - sp * cyp + (x1 + x2) / 2
    cy = sp * cxp + cp * cyp + (y1 + y2) / 2

    def ang(ux, uy, vx, vy):
        a = math.atan2(ux * vy - uy * vx, ux * vx + uy * vy)
        return a

    t1 = ang(1, 0, (x1p - cxp) / rx, (y1p - cyp) / ry)
    dt = ang((x1p - cxp) / rx, (y1p - cyp) / ry, (-x1p - cxp) / rx, (-y1p - cyp) / ry)
    if not sweep and dt > 0:
        dt -= 2 * math.pi
    elif sweep and dt < 0:
        dt += 2 * math.pi
    segs = max(1, math.ceil(abs(dt) / (math.pi / 2) - 1e-9))
    step = dt / segs
    k = 4 / 3 * math.tan(step / 4)
    out = []
    t = t1
    for _ in range(segs):
        c1, s1, c2, s2 = math.cos(t), math.sin(t), math.cos(t + step), math.sin(t + step)
        p = lambda ex, ey: (cp * rx * ex - sp * ry * ey + cx, sp * rx * ex + cp * ry * ey + cy)
        a1 = p(c1 - k * s1, s1 + k * c1)
        a2 = p(c2 + k * s2, s2 - k * c2)
        e = p(c2, s2)
        out.append((*a1, *a2, *e))
        t += step
    return out


def _path_ops(d: str) -> list[tuple]:
    """SVG path data as ('M', x, y), ('L', x, y), ('C', ...6), ('Z',) in absolute terms."""
    ops = []
    toks = TOKEN.findall(d)
    i, cmd = 0, None
    x = y = sx = sy = 0.0
    last_c = last_q = None

    def num():
        nonlocal i
        v = float(toks[i])
        i += 1
        return v

    def flag():
        # arc flags may be written without separators: "a1 1 0 011 1"
        nonlocal i
        t = toks[i]
        if len(t) > 1 and t[0] in "01":
            toks[i] = t[1:]
            return int(t[0])
        i += 1
        return int(float(t))

    while i < len(toks):
        if re.fullmatch(r"[A-Za-z]", toks[i]):
            cmd = toks[i]
            i += 1
            if cmd in "Zz":
                ops.append(("Z",))
                x, y = sx, sy
                last_c = last_q = None
                continue
        rel = cmd.islower()
        c = cmd.upper()
        ox, oy = (x, y) if rel else (0.0, 0.0)
        if c == "M":
            x, y = num() + ox, num() + oy
            sx, sy = x, y
            ops.append(("M", x, y))
            cmd = "l" if rel else "L"
            last_c = last_q = None
        elif c == "L":
            x, y = num() + ox, num() + oy
            ops.append(("L", x, y))
            last_c = last_q = None
        elif c == "H":
            x = num() + (x if rel else 0.0)
            ops.append(("L", x, y))
            last_c = last_q = None
        elif c == "V":
            y = num() + (y if rel else 0.0)
            ops.append(("L", x, y))
            last_c = last_q = None
        elif c == "C":
            x1, y1, x2, y2 = num() + ox, num() + oy, num() + ox, num() + oy
            x, y = num() + ox, num() + oy
            ops.append(("C", x1, y1, x2, y2, x, y))
            last_c, last_q = (x2, y2), None
        elif c == "S":
            x1, y1 = (2 * x - last_c[0], 2 * y - last_c[1]) if last_c else (x, y)
            x2, y2 = num() + ox, num() + oy
            x, y = num() + ox, num() + oy
            ops.append(("C", x1, y1, x2, y2, x, y))
            last_c, last_q = (x2, y2), None
        elif c in "QT":
            if c == "Q":
                qx, qy = num() + ox, num() + oy
            else:
                qx, qy = (2 * x - last_q[0], 2 * y - last_q[1]) if last_q else (x, y)
            ex, ey = num() + ox, num() + oy
            ops.append(("C", x + 2 / 3 * (qx - x), y + 2 / 3 * (qy - y),
                        ex + 2 / 3 * (qx - ex), ey + 2 / 3 * (qy - ey), ex, ey))
            x, y = ex, ey
            last_q, last_c = (qx, qy), None
        elif c == "A":
            rx, ry, phi = num(), num(), num()
            large, sweep = flag(), flag()
            ex, ey = num() + ox, num() + oy
            for seg in _arc(x, y, rx, ry, phi, large, sweep, ex, ey):
                ops.append(("C", *seg))
            x, y = ex, ey
            last_c = last_q = None
        else:
            i += 1
    return ops


def _ellipse(cx, cy, rx, ry):
    k = 0.5522847498
    return [("M", cx + rx, cy),
            ("C", cx + rx, cy + k * ry, cx + k * rx, cy + ry, cx, cy + ry),
            ("C", cx - k * rx, cy + ry, cx - rx, cy + k * ry, cx - rx, cy),
            ("C", cx - rx, cy - k * ry, cx - k * rx, cy - ry, cx, cy - ry),
            ("C", cx + k * rx, cy - ry, cx + rx, cy - k * ry, cx + rx, cy), ("Z",)]


def _f(a, key, default=0.0):
    v = a.get(key)
    return float(NUM.match(str(v)).group()) if v is not None and NUM.match(str(v)) else default


def ops_for(nodes) -> list[tuple]:
    ops = []
    for tag, a in nodes:
        if tag == "path":
            ops += _path_ops(a.get("d", ""))
        elif tag == "circle":
            r = _f(a, "r")
            ops += _ellipse(_f(a, "cx"), _f(a, "cy"), r, r)
        elif tag == "ellipse":
            ops += _ellipse(_f(a, "cx"), _f(a, "cy"), _f(a, "rx"), _f(a, "ry"))
        elif tag == "line":
            ops += [("M", _f(a, "x1"), _f(a, "y1")), ("L", _f(a, "x2"), _f(a, "y2"))]
        elif tag in ("polyline", "polygon"):
            pts = [float(v) for v in NUM.findall(a.get("points", ""))]
            if len(pts) >= 4:
                ops.append(("M", pts[0], pts[1]))
                ops += [("L", pts[k], pts[k + 1]) for k in range(2, len(pts) - 1, 2)]
                if tag == "polygon":
                    ops.append(("Z",))
        elif tag == "rect":
            x, y, w, h = _f(a, "x"), _f(a, "y"), _f(a, "width"), _f(a, "height")
            rx = _f(a, "rx", _f(a, "ry", 0.0))
            ry = _f(a, "ry", rx)
            rx, ry = min(rx, w / 2), min(ry, h / 2)
            if rx <= 0:
                ops += [("M", x, y), ("L", x + w, y), ("L", x + w, y + h), ("L", x, y + h), ("Z",)]
            else:
                k = 0.5522847498
                ops += [("M", x + rx, y), ("L", x + w - rx, y),
                        ("C", x + w - rx + k * rx, y, x + w, y + ry - k * ry, x + w, y + ry),
                        ("L", x + w, y + h - ry),
                        ("C", x + w, y + h - ry + k * ry, x + w - rx + k * rx, y + h, x + w - rx, y + h),
                        ("L", x + rx, y + h),
                        ("C", x + rx - k * rx, y + h, x, y + h - ry + k * ry, x, y + h - ry),
                        ("L", x, y + ry),
                        ("C", x, y + ry - k * ry, x + rx - k * rx, y, x + rx, y), ("Z",)]
    return ops


# ------------------------------------------------------------------ drawing

def _pt(tag, x, y):
    el = OxmlElement("a:pt")
    el.set("x", str(int(round(x * UNIT))))
    el.set("y", str(int(round(y * UNIT))))
    return el


def add(slide, name: str, x, y, size, colour: str, *, extra: str | None = None, weight: float = 2.0):
    """Draw icon `name` in a `size` square at (x, y); None when there is no such icon."""
    nodes = lookup(name, extra)
    if nodes is None:
        return None
    from .slide_layouts import _plain
    shp = _plain(slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, int(x), int(y), int(size), int(size)))
    shp.name = f"Icon {name}"
    spPr = shp._element.spPr
    prst = spPr.find(qn("a:prstGeom"))
    geom = OxmlElement("a:custGeom")
    for tag in ("a:avLst", "a:gdLst", "a:ahLst", "a:cxnLst"):
        geom.append(OxmlElement(tag))
    rect = OxmlElement("a:rect")
    for k, v in (("l", "l"), ("t", "t"), ("r", "r"), ("b", "b")):
        rect.set(k, v)
    geom.append(rect)
    lst = OxmlElement("a:pathLst")
    path = OxmlElement("a:path")
    path.set("w", str(int(GRID * UNIT)))
    path.set("h", str(int(GRID * UNIT)))
    path.set("fill", "none")
    for op in ops_for(nodes):
        if op[0] == "M":
            el = OxmlElement("a:moveTo")
            el.append(_pt("a:pt", op[1], op[2]))
        elif op[0] == "L":
            el = OxmlElement("a:lnTo")
            el.append(_pt("a:pt", op[1], op[2]))
        elif op[0] == "C":
            el = OxmlElement("a:cubicBezTo")
            for k in range(3):
                el.append(_pt("a:pt", op[1 + 2 * k], op[2 + 2 * k]))
        else:
            el = OxmlElement("a:close")
        path.append(el)
    lst.append(path)
    geom.append(lst)
    prst.addprevious(geom)
    spPr.remove(prst)
    shp.fill.background()
    from .slides import _rgb
    shp.line.color.rgb = _rgb(colour)
    shp.line.width = Emu(int(size * weight / GRID))
    ln = shp.line._get_or_add_ln()
    ln.set("cap", "rnd")
    for old in ln.findall(qn("a:round")):
        ln.remove(old)
    ln.append(OxmlElement("a:round"))
    return shp
